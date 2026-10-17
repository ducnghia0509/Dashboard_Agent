# -*- coding: utf-8 -*-
"""Khoá các bẫy của `derive_hcns_bangluong` — 'Bảng lương all' của HCNS.

Chạy: python scripts/test_hcns_bangluong.py

Mỗi ca dưới đây đều làm số SAI IM LẶNG nếu vỡ (đo thật 17/10/2026):
  · tên 'Thực lĩnh' / 'Thuế' / 'BH' có ở CẢ phần chi tiết lẫn khối tổng hợp cuối bảng — lấy
    nhầm cột đầu là đọc số của cột khác;
  · cột 'Thực lĩnh' từng người đã trừ 'Giữ lương/lương âm', Bctienluong thì chưa;
  · 'Khối Hưng Thịnh' giấu khối thật ở cột phòng;
  · tên công ty đổi theo tháng ('Vinfast Thịnh Cường' -> 'Chi nhánh Vinfast - …');
  · bản 'Bảng lương all T9' là MẪU RỖNG — phải bị loại, không được ghi 0 đè lên kỳ.
"""
import datetime as dt
import os
import sys
import tempfile

import openpyxl

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)

import derive_hcns_bangluong as D  # noqa: E402

TONG = ["Lương CB", "Lương ý thức chất lượng công việc", "Tổng phụ cấp", "Thưởng thành tích",
        "Thưởng KPI", "Thưởng TN&Doanh số", "Thưởng chuyên cần", "Các khoản phúc lợi",
        "Truy thu truy lĩnh", "các khoản giảm trừ", "BH", "Công đoàn", "Thuế", "Hoàn thuế 5%",
        "Thực lĩnh", "Giữ lương/lương âm"]
CTY = ["Vinfast Thịnh Cường", "Công ty Cổ phần Thịnh Cường", "Công ty Cổ phần Global AI",
       "Công ty Cổ phần An An's Garden", "Công ty CP Công nghệ và dịch vụ Xanh Vĩnh Phúc",
       "Công ty TNHH Xuất nhập khẩu và khai thác Hưng Thịnh"]


def _wb(path, nhan_vien, dong_td=9):
    wb = openpyxl.Workbook()
    wb.active.title = "TH"
    ws = wb.create_sheet("Bảng lương ")
    ws.cell(1, 1, "BẢNG LƯƠNG")
    # Phần chi tiết CŨNG có 'Thực lĩnh', 'Thuế', 'BH' — mồi nhử cột trùng tên.
    td = ["STT", "Công ty", "Khối ", "Phòng/Bộ phận", "MÃ NV", "HỌ VÀ TÊN", "Ngày nghỉ việc",
          "BH", "Thuế", "Thực lĩnh", "Ban"] + TONG
    for j, h in enumerate(td, 1):
        ws.cell(dong_td, j, h)
    for i, nv in enumerate(nhan_vien, dong_td + 1):
        ws.cell(i, 1, i)
        ws.cell(i, 2, nv["cty"])
        ws.cell(i, 3, nv["khoi"])
        ws.cell(i, 4, nv.get("phong", "P"))
        ws.cell(i, 5, nv.get("ma", f"{i:06d}"))
        ws.cell(i, 7, nv.get("nghi", 0))
        ws.cell(i, 8, 999)
        ws.cell(i, 9, 999)
        ws.cell(i, 10, 999)
        g = {"Lương CB": 10e6, "Thưởng KPI": 1e6, "Thưởng TN&Doanh số": 5e6, "BH": 1.05e6,
             "Thuế": 0.5e6, "Giữ lương/lương âm": 2e6, **nv.get("tien", {})}
        for j, h in enumerate(TONG, 12):
            ws.cell(i, j, g.get(h, 0))
        ws.cell(i, 12 + TONG.index("Thực lĩnh"), -1)       # cột từng người: KHÔNG được dùng
    wb.save(path)


def _mot_nguoi_moi_cty(**them):
    return [{"cty": c, "khoi": "Khối Hỗ trợ", **them} for c in CTY]


def test_lay_cot_tong_hop_cuoi_bang_va_tu_tinh_thuc_linh():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "all.xlsx")
        _wb(p, _mot_nguoi_moi_cty())
        r = D.doc_file(p, "2026-07")
        assert r and not r.get("loi"), r
        x = r["recs"][0]
        # CP = 10 + 1 + 5 = 16tr; thưởng = KPI 1tr (TN&DS không tính); TL = 16 - 1.05 - 0.5
        assert x["payload"]["tongCP"] == 16e6
        assert x["payload"]["thuong"] == 1e6
        assert x["amount"] == 15e6
        assert abs(x["amount2"] - 14.45e6) < 1e-3, x["amount2"]     # chưa trừ giữ lương 2tr
        assert abs(x["payload"]["bhxh"] - 3.2e6) < 1e-3              # 1,05tr × 32/10,5


def test_ten_cong_ty_quy_ve_ten_dashboard():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "all.xlsx")
        _wb(p, _mot_nguoi_moi_cty())
        ten = {r["dim3"] for r in D.doc_file(p, "2026-07")["recs"]}
        assert "Chi nhánh Vinfast - Công ty Thịnh Cường" in ten
        assert "Công ty Xanh Vĩnh Phúc" in ten
        assert "Vinfast Thịnh Cường" not in ten


def test_khoi_mang_ten_cong_ty_lay_khoi_that_o_cot_phong():
    nv = _mot_nguoi_moi_cty()
    nv[-1] = {"cty": CTY[-1], "khoi": "Khối Hưng Thịnh", "phong": "Khối kinh doanh xe tải"}
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "all.xlsx")
        _wb(p, nv)
        ht = [r for r in D.doc_file(p, "2026-08")["recs"] if "Hưng Thịnh" in r["dim3"]]
        assert ht[0]["dim1"] == "Khối kinh doanh xe tải", ht


def test_nghi_viec_trong_thang_khong_vao_lv_cuoi_ky():
    nv = _mot_nguoi_moi_cty() + [
        {"cty": CTY[1], "khoi": "Khối Hỗ trợ", "nghi": dt.datetime(2026, 7, 15)},
        {"cty": CTY[1], "khoi": "Khối Hỗ trợ", "nghi": dt.datetime(2026, 8, 3)}]
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "all.xlsx")
        _wb(p, nv)
        tc = [r for r in D.doc_file(p, "2026-07")["recs"] if r["dim3"] == CTY[1]][0]
        assert tc["payload"]["soNguoiBangLuong"] == 3
        assert tc["payload"]["lvCuoiKy"] == 2          # người nghỉ 15/07 không còn cuối tháng 7


def test_tieu_de_troi_xuong_dong_16():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "all.xlsx")
        _wb(p, _mot_nguoi_moi_cty(), dong_td=16)
        assert D.doc_file(p, "2026-09")["dong_tieu_de"] == 16


def test_mau_rong_bi_loai_va_chon_ban_co_so():
    with tempfile.TemporaryDirectory() as d:
        rong, that = os.path.join(d, "Bảng lương all T9.xlsx"), os.path.join(d, "theo SĐTC.xlsx")
        _wb(rong, [{"cty": c, "khoi": "K", "ma": ""} for c in CTY])   # có khuôn, không có ai
        _wb(that, _mot_nguoi_moi_cty())
        p, r, bo = D.chon("2026-09", [rong, that])
        assert p == that, (p, bo)
        assert any("mẫu rỗng" in b["ly_do"] for b in bo), bo


def test_bang_luong_mot_cong_ty_khong_phai_all():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "F.9.TC.xlsx")
        _wb(p, [{"cty": CTY[1], "khoi": "K"}] * 3)
        assert D.doc_file(p, "2026-07") is None


def test_ky_lay_tu_ten_thu_muc():
    assert D.ky_cua_thu_muc("BANG LUONG NAM 2026 - Bang luong thang 07.2026 - Luong thang 7 "
                            "(theo SĐTC mới)") == "2026-07"
    assert D.ky_cua_thu_muc("BANG LUONG NAM 2026 - Bang luong thang 06.2026 - Bang luong "
                            "T6.2026_SĐTCM") == "2026-06"
    assert D.ky_cua_thu_muc("BANG LUONG NAM 2026 - Thuong Nam 2026 - Thuong PD thang 02.2026") is None


if __name__ == "__main__":
    n = 0
    for ten, f in list(globals().items()):
        if ten.startswith("test_") and callable(f):
            f()
            n += 1
            print("OK", ten)
    print(f"{n} test đạt")
