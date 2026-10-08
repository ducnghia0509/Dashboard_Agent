# -*- coding: utf-8 -*-
"""KẾ HOẠCH DỰ THU THEO NGÀY -> report_type `THUCHI_KH` (cột "Kế hoạch" của mục Thu trong kỳ, màn Dòng tiền).

SPEC (user 08/10/2026, "Dòng tiền dự thu đi theo pháp nhân"): mỗi đơn vị một file
`<thư mục đơn vị>\\CONGNOPHAITHUNGAY\\B.x.<pháp nhân>.TCKT.M.<YYYYMM>_Kehoachthu<đơn vị>.xlsx`
(Xanh VP: `XANHVINHPHUC\\KEHOACHPHAITHUNGAY`), sheet KHT_*, DÒNG 8 "Thu từ hoạt động kinh doanh",
một cột mỗi ngày trong tháng. Cost center:
  Showroom  cộng các dòng cost center của 3 kênh B2C / B2B / GF
  XDV       cộng các dòng cost center của 2 loại hình Thu bảo hiểm / Thu khách lẻ
  Xanh VP   Depot Phú Thọ dòng 10 · Tuyên Quang dòng 22 · Vĩnh Phúc dòng 34
  An Taxi   Depot Sơn Tây dòng 10 · Thái Nguyên dòng 34
  còn lại   chỉ dòng 8
GlobalAI: spec ghi "chưa có" (file T09 trống) -> không khai.

ĐỌC THEO NHÃN, KHÔNG THEO CHỮ CỘT / SỐ DÒNG: file T10 của An Taxi và An KS chèn thêm cột "Tuần 5"
nên cột ngày dời từ P->AT sang Q->AU — số ngày đọc từ chính tiêu đề "Ngày 01".."Ngày 31". Dòng tổng
= dòng STT "1" đầu tiên dưới tiêu đề (Xe tải T10 đổi nhãn thành "Tổng dòng thu"), khối kết thúc ở
STT "2" (Thu từ hoạt động tài chính). Nhóm = STT "1.x", dòng cost center = nhãn cột C bắt đầu "+".

KHÔNG DÙNG CỘT TỔNG (K): An Taxi/An KS T10 cột K chỉ cộng tuần 1..4, thiếu tuần 5 (An KS: K 49,5 tr,
Σ ngày 60 tr). Kỳ lấy theo TÊN FILE: tiêu đề "Kỳ" trong file sai ở Trạm sạc T10 (ghi tháng 9) và
Dự án T09 (ghi tháng 10).

PHÁP NHÂN: dòng tiền tách theo pháp nhân như THUCHI thực tế. Showroom Uông Bí thuộc VFQN theo danh
mục cost center (org_catalog) -> dòng kế hoạch của nó gắn VFQN, khớp cột thực tế của VFQN.

COST CENTER NẰM TRONG payload, KHÔNG ở cột cost_center: `repository.flow_sum` gom theo file — file
có dòng KHÔNG cost center thì CHỈ lấy các dòng đó (coi là dòng tổng). Showroom có '+ Vinfast Hạ Long 1'
chưa có mã -> để ở cột là tổng kế hoạch Showroom chỉ còn đúng dòng đó. THUCHI thực tế cũng không có
cost center nên lọc theo cột không có nghĩa ở màn này.

Σ dòng cost center phải bằng dòng tổng từng ngày; phần chênh (nếu có) ghi một dòng cost center rỗng
`chenh_lech_tong` để tổng đơn vị luôn đúng dòng 8.

  DATABASE_URL=... .venv/bin/python scripts/derive_kehoach_thu_ngay.py [--period 2026-10 ...] [--write]
"""
import argparse
import calendar
import datetime as dt
import glob
import json
import os
import re
import sys
import unicodedata

import openpyxl
import psycopg

AGENT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_URL = (os.environ.get("DATABASE_URL") or os.environ.get("TC_DATABASE_URL")
          or "postgresql://tc:tc@localhost:5433/tc_dashboard")
RECEIVED_DIR = os.environ.get("RECEIVED_DIR") or os.path.normpath(
    os.path.join(AGENT, "..", "Connect_VPS", "received_reports"))
RT = "THUCHI_KH"
KM = "A. Thu từ hoạt động kinh doanh"

# (mã đơn vị, thư mục nhận, hậu tố tên file, sheet, pháp nhân, khối, chế độ cost center, khối catalog)
_UNITS = [
    ("SRVF", "SRVF", "KDXE", "KHT_KDXE", "TC", "Khối KD Vinfast - Showroom", "nhom", "1"),
    ("XDV", "XDV", "XDV", "KHT_KDXDV", "TC", "Khối KD Vinfast - XDV", "nhom", "2"),
    ("TRAMSAC", "TRAMSAC", "Tramsac", "KHT_KDTRAMSAC", "TC", "Khối KD Trạm sạc Vgreen", None, None),
    ("HO", "HO", "HO", "KHT_KDGA", "TC", "Khối hỗ trợ tập đoàn", None, None),
    ("DUAN", "DUAN", "DuAn", "KHT_KDDUAN", "TC", "Khối KD Dự án", None, None),
    ("HUNGTHINH", "HUNGTHINH", "Xetai", "KHT_KDXETai", "HT", "Khối KD Xe tải", None, None),
    # Thư mục nhận: tới chiều 08/10/2026 máy Local xếp file XANHVINHPHUC\KEHOACHPHAITHUNGAY vào company
    # 'CONGNOPHAITHU' (cùng chỗ bộ file B.9 cũ của TCGROUP\PHAITHUTAPDOAN), sau đó phân loại lại thành
    # XANHVINHPHUC/kehoachthungay -> tìm cả hai, lấy bản mới nhất. Hậu tố 'XanhVP' chỉ khớp file Xanh VP.
    ("XVP", ("XANHVINHPHUC", "CONGNOPHAITHU"), "XanhVP", "KHT_KDXanhTaxi", "XVP", "Khối KD Vận tải Taxi Xanh", "depot", "6"),
    ("ANTAXI", "ANTAXI", "AnTaxi", "KHT_KDAnTaxi", "AAG", "Khối KD Dịch vụ An Taxi", "depot", "7"),
    ("ANKHACHSAN", "ANKHACHSAN", "AnKS", "KHT_KDHO", "AAG", "Khối KD Dịch vụ An KS", None, None),
]
# Nhãn dòng không trùng tên cost center trong danh mục.
_CC_ALIAS = {"kd b2b": "B2B_SR"}
_RX_NGAY = re.compile(r"^\s*ng[aà]y\s*(\d{1,2})\s*$", re.I)


def _nd(s):
    s = str(s or "").strip().lower().replace("đ", "d")
    return "".join(ch for ch in unicodedata.normalize("NFD", s) if unicodedata.category(ch) != "Mn")


def _ten_cc(s):
    """'+ Vinfast Xuân Mai (OP: T4)' / 'Vinfast Xuân Mai' -> 'xuan mai'."""
    s = re.sub(r"\(.*?\)", "", _nd(s)).lstrip("+-").strip()
    s = re.sub(r"^vinfast\s+", "", s)
    return re.sub(r"\s+", " ", s).strip()


def _catalog():
    with open(os.path.join(AGENT, "servers", "common", "org_catalog.json"), encoding="utf-8") as f:
        return json.load(f).get("cost_centers", [])


def _cc_map(khoi_cat):
    out = {}
    for c in _catalog():
        if c.get("khoi") == khoi_cat and not re.search(r"_\d+$", c["ma"]):   # bỏ mã '(61)' của XVP
            out.setdefault(_ten_cc(c["ten"]), (c["ma"], c.get("cong_ty")))
    return out


def _stt(v):
    return str(v).strip() if v is not None else ""


def _num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def doc_file(path, unit, period):
    """-> (records, canh_bao). record = dict(ngay, amount_vnd, cong_ty, cost_center, kenh, ten_dong)."""
    ma, _, _, sheet, cong_ty, khoi, mode, khoi_cat = unit
    y, m = int(period[:4]), int(period[5:7])
    so_ngay = calendar.monthrange(y, m)[1]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return None, [f"không có sheet {sheet}"]
        rows = [list(r) for r in wb[sheet].iter_rows(min_row=1, max_row=200, max_col=70, values_only=True)]
    finally:
        wb.close()
    canh_bao = []
    hi = next((i for i, r in enumerate(rows[:15])
               if sum(1 for c in r if isinstance(c, str) and _RX_NGAY.match(c)) >= 28), None)
    if hi is None:
        return None, ["không thấy dòng tiêu đề 'Ngày 01..'"]
    cot_ngay = {j: int(_RX_NGAY.match(c).group(1)) for j, c in enumerate(rows[hi])
                if isinstance(c, str) and _RX_NGAY.match(c)}
    t0 = next((i for i in range(hi + 1, len(rows)) if _stt(rows[i][0]) == "1"), None)
    if t0 is None:
        return None, ["không thấy dòng STT 1 (Thu từ hoạt động kinh doanh)"]
    t1 = next((i for i in range(t0 + 1, len(rows)) if re.fullmatch(r"\d+", _stt(rows[i][0]))), len(rows))

    def ngay_cua(r):
        out = {}
        for j, d in cot_ngay.items():
            v = _num(r[j]) if j < len(r) else 0.0
            if not v:
                continue
            if d > so_ngay:
                # Mẫu file luôn có 31 cột ngày; T09 kế toán vẫn điền 'Ngày 31' (Showroom 5,03 tỷ, Xe tải
                # 13,33 tỷ). Bỏ đi thì kế hoạch tháng hụt so với file -> dồn về ngày cuối tháng.
                canh_bao.append(f"cột 'Ngày {d}' có số {v:,.0f} nhưng tháng chỉ có {so_ngay} ngày"
                                f" — dồn vào ngày {so_ngay}")
                d = so_ngay
            out[d] = out.get(d, 0.0) + v
        return out

    tong = ngay_cua(rows[t0])
    la = []          # (ngay_dict, cost_center, cong_ty, kenh, ten_dong)
    if mode:
        ccm = _cc_map(khoi_cat)
        nhom = None
        for i in range(t0 + 1, t1):
            r = rows[i]
            stt, nhan = _stt(r[0]), (r[2] if len(r) > 2 else None)
            if re.fullmatch(r"1\.\d+", stt):
                # '- GF (bán xe cũ Vinfast)' -> 'GF'
                nhom = re.split(r"\s*\(", str(nhan or "").strip(" -"))[0] or None
                if mode == "depot":
                    vals = ngay_cua(r)
                    if vals:
                        hit = ccm.get(_ten_cc(nhan)) if nhan else None
                        la.append((vals, hit[0] if hit else None, cong_ty, None, str(nhan or f"dòng {i + 1}")))
                        if not hit:
                            canh_bao.append(f"dòng {i + 1} '{nhan}' không khớp cost center nào")
                continue
            if mode == "nhom" and isinstance(nhan, str) and nhan.strip().startswith("+"):
                vals = ngay_cua(r)
                if not vals:
                    continue
                key = _ten_cc(nhan)
                ma_cc = _CC_ALIAS.get(key)
                hit = (ma_cc, next((c.get("cong_ty") for c in _catalog() if c["ma"] == ma_cc), None)) \
                    if ma_cc else ccm.get(key)
                if not hit:
                    canh_bao.append(f"dòng {i + 1} '{nhan.strip()}' không khớp cost center nào")
                la.append((vals, hit[0] if hit else None, (hit[1] if hit and hit[1] else cong_ty),
                           nhom, nhan.strip()))
    if not la:
        la = [(tong, None, cong_ty, None, str(rows[t0][1] or "dòng tổng"))]
    else:   # Σ cost center phải khớp dòng tổng từng ngày
        for d in sorted(set(tong) | {d for v, *_ in la for d in v}):
            chenh = tong.get(d, 0.0) - sum(v.get(d, 0.0) for v, *_ in la)
            if abs(chenh) >= 1:
                la.append(({d: chenh}, None, cong_ty, None, "chenh_lech_tong"))
                canh_bao.append(f"ngày {d}: Σ cost center lệch dòng tổng {chenh:,.0f} đ")
    recs = []
    for vals, cc, cty, kenh, ten in la:
        for d, v in vals.items():
            recs.append({"ngay": dt.date(y, m, d).isoformat(), "vnd": v, "cong_ty": cty,
                         "cost_center": cc, "kenh": kenh, "ten_dong": ten})
    return recs, canh_bao


def _thu_muc(unit):
    return unit[1] if isinstance(unit[1], tuple) else (unit[1],)


def _file_moi_nhat(unit, period):
    """Bản mới nhất của đơn vị/kỳ. Cùng một file có thể nằm ở 2 thư mục con (baocaotuoino trước
    08/10/2026, kehoachthungay sau khi máy Local phân loại lại) -> so mtime."""
    hau_to = unit[2]
    rx = re.compile(r"\.M\.%s%s_Kehoachthu%s\.xlsx$" % (period[:4], period[5:7], re.escape(hau_to)), re.I)
    cands = [p for f in _thu_muc(unit)
             for p in glob.glob(os.path.join(RECEIVED_DIR, f, "**", "*.xlsx"), recursive=True)
             if rx.search(os.path.basename(p))]
    return max(cands, key=os.path.getmtime) if cands else None


def _ky_co_file():
    out = set()
    for u in _UNITS:
        for p in (p for f in _thu_muc(u)
                  for p in glob.glob(os.path.join(RECEIVED_DIR, f, "**", "*_Kehoachthu%s.xlsx" % u[2]),
                                     recursive=True)):
            m = re.search(r"\.M\.(\d{4})(\d{2})_", os.path.basename(p))
            if m:
                out.add(f"{m.group(1)}-{m.group(2)}")
    return sorted(out)


def chay(periods, write=False):
    kq = []
    conn = psycopg.connect(DB_URL) if write else None
    try:
        cur = conn.cursor() if conn else None
        for period in periods:
            ds = None
            if cur:
                cur.execute("SELECT dataset_id FROM raw_rows WHERE period_month=%s "
                            "ORDER BY (report_type='THUCHI') DESC LIMIT 1", (period,))
                r = cur.fetchone()
                ds = r[0] if r else None
            for u in _UNITS:
                path = _file_moi_nhat(u, period)
                if not path:
                    kq.append({"ky": period, "don_vi": u[0], "ghi": False, "ly_do": "chưa có file"})
                    continue
                recs, cb = doc_file(path, u, period)
                item = {"ky": period, "don_vi": u[0], "file": os.path.basename(path), "canh_bao": cb}
                if recs is None:
                    item.update(ghi=False, ly_do="; ".join(cb))
                    kq.append(item)
                    continue
                item.update(dong=len(recs), tong_ty=round(sum(r["vnd"] for r in recs) * 1e-9, 9),
                            theo_phap_nhan={c: round(sum(r["vnd"] for r in recs if r["cong_ty"] == c) * 1e-9, 9)
                                            for c in sorted({r["cong_ty"] for r in recs})})
                if write:
                    if not ds:
                        item.update(ghi=False, ly_do=f"chưa có dataset kỳ {period}")
                        kq.append(item)
                        continue
                    src = f"{os.path.relpath(path, RECEIVED_DIR).split(os.sep)[0]}::{os.path.basename(path)}"
                    cur.execute("DELETE FROM raw_rows WHERE report_type=%s AND period_month=%s AND dim3=%s",
                                (RT, period, u[0]))
                    for k, r in enumerate(recs):
                        cur.execute(
                            "INSERT INTO raw_rows (dataset_id, report_type, row_index, ngay, cong_ty, khoi, "
                            "cost_center, period_month, amount, amount2, dim1, dim2, dim3, payload, source_file) "
                            "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,NULL,%s,NULL,%s,%s,%s)",
                            (ds, RT, 9000000 + k, r["ngay"], r["cong_ty"], u[5], None, period,
                             r["vnd"] * 1e-9, KM, u[0],
                             json.dumps({"unit": "ty", "cost_center": r["cost_center"], "kenh": r["kenh"],
                                         "ten_dong": r["ten_dong"]}, ensure_ascii=False), src))
                    conn.commit()
                item["ghi"] = write
                kq.append(item)
    finally:
        if conn:
            conn.close()
    return kq


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--period", action="append", help="YYYY-MM (lặp lại được); bỏ trống = mọi kỳ có file")
    ap.add_argument("--write", action="store_true", help="ghi DB (mặc định dry-run)")
    a = ap.parse_args()
    print(json.dumps(chay(a.period or _ky_co_file(), a.write), ensure_ascii=False, indent=1))


if __name__ == "__main__":
    sys.exit(main())
