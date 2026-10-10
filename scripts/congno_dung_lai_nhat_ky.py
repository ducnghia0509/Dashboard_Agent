# -*- coding: utf-8 -*-
"""Dựng lại sổ công nợ TK131/TK331 THEO ĐỐI TƯỢNG khi sheet công nợ của file bị ĐỨNG YÊN.

CA THẬT (10/10/2026, khối HO): sheet 'PTHU'/'PTRA' trong file BCTC tháng HO là mẫu SUMIFS theo khoảng
ngày ở ô H2:I2 (PTRA: G2:H2). Từ T01/2026 hai ô đó vẫn ghi 01/04/2025–30/04/2025 trong khi sheet
'DATA' (nhật ký) chỉ có chứng từ 2026 -> phát sinh = 0, cuối kỳ = cột 'Dư đầu kỳ' tĩnh cũ. Mọi tháng
T01–T08 ra cùng 240,28 tỷ phải thu / 18,28 tỷ trả trước, trong khi CĐKT cùng file ghi 131 = 0,49 tỷ.

CÁCH DỰNG (tất định, không đoán):
  số dư từng đối tượng = số dư CUỐI KỲ ở sheet công nợ của FILE GẦN NHẤT CÓ SHEET ĐÚNG KỲ
                         (HO: T12/2025, sheet ghi 01/12/2025–31/12/2025)
                       + phát sinh sheet 'DATA' của chính file đang nạp (TKGHINO/TKGHICO 131x/331x,
                         'Đối tượng nợ'/'Đối tượng có'), cộng dồn theo cột 'Tháng'.
  Đối chiếu từng tháng với dòng TK cha (1311/3311…) của sheet 'CĐPS' cùng file: phần lệch (số dư đầu
  năm sổ 2026 khác sheet T12 — HO: 131 lệch 0,109743 tỷ, 331 lệch ~1,664 tỷ, KHÔNG đổi qua các tháng,
  tức phát sinh khớp tới đồng) ghi thành MỘT dòng 'Chênh lệch số dư đầu năm chưa phân bổ theo đối
  tượng' cho từng TK -> tổng luôn bằng CĐPS/CĐKT của đúng tháng đó. Không phân bổ phần lệch cho đối
  tượng nào (bịa số).

Trả về cùng hình dạng `records`/`adv_recs` mà `agent_cli._derive_congno` dựng từ sheet thật, để đi
chung một đường ghi (template 05/06 -> PTHU/PTRA, dư ngược chiều -> PTHU_ADV/PTRA_ADV).
"""
import collections
import datetime as dt
import glob
import os
import re

TEN_CHUA_PHAN_BO = "Chênh lệch số dư đầu năm chưa phân bổ theo đối tượng"
_RE_KY_FILE = re.compile(r"\.M\.(\d{4})(\d{2})\.", re.I)


def _so(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def ky_dau_sheet(rows):
    """(năm, tháng) mà sheet công nợ TỰ KHAI ở các ô NGÀY đầu sheet (ô 'Đến ngày' = ngày lớn nhất
    trong 6 dòng đầu), hoặc None nếu sheet không khai ô ngày nào."""
    ngay = [c for r in rows[:6] for c in r if isinstance(c, (dt.datetime, dt.date))]
    if not ngay:
        return None
    m = max(ngay)
    return (m.year, m.month)


def _mo(path):
    # openpyxl thẳng (không qua be_bridge): module chỉ đọc .xlsx BCTC tháng, và phải chạy được cả ở
    # nơi không cấu hình BACKEND_PATH (kiểm thử, worktree).
    import openpyxl
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def _sheet_cong_no(wb, canonical_kind):
    want = {"TK131": ("pthu", "131"), "TK331": ("ptra", "331")}[canonical_kind]
    for n in wb.sheetnames:
        if n.strip().lower() in want:
            return n
    return None


def _so_du_mo_dau(file_path, canonical_kind, period):
    """{(TK4, mã đối tượng): số dư ròng VND}, {mã: tên}, kỳ của file gốc — từ file GẦN NHẤT (kỳ <
    period) cùng thư mục có sheet công nợ khai ĐÚNG kỳ của nó. None nếu không có file nào."""
    tk_goc = "131" if canonical_kind == "TK131" else "331"
    ung_vien = []
    for p in glob.glob(os.path.join(os.path.dirname(file_path), "*.xls*")):
        m = _RE_KY_FILE.search(os.path.basename(p))
        if m and f"{m.group(1)}-{m.group(2)}" < period and not os.path.basename(p).startswith("~$"):
            ung_vien.append((f"{m.group(1)}-{m.group(2)}", p))
    for ky, p in sorted(ung_vien, reverse=True):
        try:
            wb = _mo(p)
        except Exception:  # noqa: BLE001
            continue
        try:
            sh = _sheet_cong_no(wb, canonical_kind)
            if not sh:
                continue
            rows = [list(r) for r in wb[sh].iter_rows(max_col=14, values_only=True)]
        finally:
            wb.close()
        kd = ky_dau_sheet(rows)
        if not kd or f"{kd[0]}-{kd[1]:02d}" != ky:
            continue                      # sheet của file này cũng đứng yên -> lùi tiếp
        # Cột: B=TK, C=mã, D=tên; cuối kỳ Nợ/Có ở 2 cột "DƯ CUỐI KỲ" (dò theo dòng tiêu đề).
        hdr = next((i for i, r in enumerate(rows[:10])
                    if any(isinstance(c, str) and "cuối kỳ" in c.lower() for c in r)), None)
        if hdr is None:
            continue
        j_ck = next(j for j, c in enumerate(rows[hdr]) if isinstance(c, str) and "cuối kỳ" in c.lower())
        du, ten = collections.defaultdict(float), {}
        for r in rows[hdr + 2:]:
            tk, ma = str(r[1] or "").strip(), str(r[2] or "").strip()
            if not tk.startswith(tk_goc) or not ma:
                continue
            du[(tk[:4], ma)] += _so(r[j_ck]) - _so(r[j_ck + 1]) if tk_goc == "131" \
                else _so(r[j_ck + 1]) - _so(r[j_ck])
            ten[ma] = str(r[3] or ma).strip()
        return du, ten, ky
    return None


def _cdps_tk_cha(wb, tk_goc):
    """{TK4: (dư đầu ròng, dư cuối ròng) VND} từ dòng TK CHA 4 số (1311, 3311…) của sheet 'CĐPS'.
    Ròng theo chiều chính của TK (131: Nợ−Có; 331: Có−Nợ)."""
    sh = next((n for n in wb.sheetnames if n.strip().upper() in ("CĐPS", "CDPS")), None)
    if not sh:
        return None
    out = {}
    for r in wb[sh].iter_rows(min_row=8, max_col=16, values_only=True):
        tk = str(r[4] or "").strip()
        if len(tk) == 4 and tk.startswith(tk_goc):
            dau = _so(r[9]) - _so(r[10])
            cuoi = _so(r[13]) - _so(r[14])
            out[tk] = (dau, cuoi) if tk_goc == "131" else (-dau, -cuoi)
    return out


def dung_lai(file_path, canonical_kind, period):
    """-> {"records": [...], "adv_recs": [...], "goc": kỳ file gốc, "chua_phan_bo": {TK: tỷ}} hoặc
    {"error": ...}. Số đã quy về TỶ, cùng khoá cột với `_derive_congno`."""
    tk_goc = "131" if canonical_kind == "TK131" else "331"
    nam, thang = int(period[:4]), int(period[5:7])
    mo = _so_du_mo_dau(file_path, canonical_kind, period)
    if not mo:
        return {"error": "không có file kỳ trước nào có sổ công nợ đúng kỳ để làm số dư đầu"}
    du, ten, ky_goc = mo
    nam_goc, thang_goc = int(ky_goc[:4]), int(ky_goc[5:7])
    if nam_goc not in (nam, nam - 1):
        return {"error": f"file gốc {ky_goc} cách quá xa kỳ {period} (nhật ký chỉ có năm {nam})"}
    tu_thang = thang_goc + 1 if nam_goc == nam else 1

    wb = _mo(file_path)
    try:
        sh = next((n for n in wb.sheetnames if n.strip().upper() == "DATA"), None)
        if not sh:
            return {"error": "file không có sheet 'DATA' (nhật ký) để dựng phát sinh"}
        ps = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0, 0.0]))  # tháng->k->[N,C]
        for r in wb[sh].iter_rows(min_row=3, max_col=18, values_only=True):
            th, tien = r[1], _so(r[8])
            if not isinstance(th, (int, float)) or not tien or not (tu_thang <= int(th) <= thang):
                continue
            for j_tk, j_ma, j_ten, phia in ((5, 14, 15, 0), (6, 16, 17, 1)):
                tk = str(r[j_tk] or "").strip()
                if not tk.startswith(tk_goc):
                    continue
                ma = str(r[j_ma] or "").strip()
                if ma in ("", "0"):
                    # Chứng từ KHÔNG ghi mã đối tượng: vế Nợ thường trống cả tên ('Bán xe audi'),
                    # vế Có lại ghi tên NGƯỜI NỘP TIỀN ('Kiều Hoàn/BP tài sản') — tách theo tên là
                    # đẻ ra phải thu ảo ở một bên, người mua trả trước ảo ở bên kia (HO 08/2026:
                    # 0,58 tỷ mỗi bên, ròng ~0). Gom một nhóm theo TK để hai vế tự triệt.
                    ma = f"KHONG_MA_{tk[:4]}"
                    ten[ma] = f"Chứng từ không ghi mã đối tượng ({tk[:4]})"
                ten.setdefault(ma, str(r[j_ten] or ma).strip())
                ps[int(th)][(tk[:4], ma)][phia] += tien
        cdps = _cdps_tk_cha(wb, tk_goc)
    finally:
        wb.close()
    if not cdps:
        return {"error": "file không có sheet 'CĐPS' để đối chiếu"}

    def _rong(nc):   # [Nợ, Có] -> biến động theo chiều chính
        return nc[0] - nc[1] if tk_goc == "131" else nc[1] - nc[0]

    for m in range(tu_thang, thang):          # cộng dồn tới ĐẦU tháng đang nạp
        for k, nc in ps[m].items():
            du[k] += _rong(nc)
    dau = dict(du)
    for k, nc in ps[thang].items():
        du[k] += _rong(nc)

    tang_i, giam_i = (0, 1) if tk_goc == "131" else (1, 0)
    records, adv, chua_pb = [], [], {}
    for k in sorted(set(dau) | set(du) | set(ps[thang])):
        d0, d1 = dau.get(k, 0.0), du.get(k, 0.0)
        nc = ps[thang].get(k, [0.0, 0.0])
        if not any(abs(x) >= 0.5 for x in (d0, d1, nc[0], nc[1])):
            continue
        records.append({"tk": k[0], "ma": k[1], "ten": ten.get(k[1], k[1]), "dau": d0, "cuoi": d1,
                        "tang": nc[tang_i], "giam": nc[giam_i]})
    for tk, (cd0, cd1) in cdps.items():
        s0 = sum(r["dau"] for r in records if r["tk"] == tk)
        s1 = sum(r["cuoi"] for r in records if r["tk"] == tk)
        if abs(cd0 - s0) >= 0.5 or abs(cd1 - s1) >= 0.5:
            records.append({"tk": tk, "ma": None, "ten": f"{TEN_CHUA_PHAN_BO} ({tk})",
                            "dau": cd0 - s0, "cuoi": cd1 - s1, "tang": 0.0, "giam": 0.0})
            chua_pb[tk] = round((cd1 - s1) * 1e-9, 9)

    t9 = lambda v: round(v * 1e-9, 9)  # noqa: E731
    out = []
    for r in records:
        out.append({"ten": r["ten"], "ma": r["ma"], "dau": t9(r["dau"]), "cuoi": t9(r["cuoi"]),
                    "tang": t9(r["tang"]), "giam": t9(r["giam"])})
        # Dư NGƯỢC chiều (131 dư Có / 331 dư Nợ) = phần âm của số dư ròng.
        if r["cuoi"] < -0.5 or r["dau"] < -0.5:
            adv.append({"ten": r["ten"], "ma": r["ma"], "cuoi": t9(max(-r["cuoi"], 0.0)),
                        "dau": t9(max(-r["dau"], 0.0))})
    return {"records": out, "adv_recs": adv, "goc": ky_goc, "chua_phan_bo": chua_pb,
            "tong_cuoi": {tk: t9(v[1]) for tk, v in cdps.items()}}
