#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Deriver: LƯƠNG THƯỞNG theo CÔNG TY × KHỐI × PHÒNG từ "Bảng lương all" của HCNS -> HCNS_LUONG_BL.

Nuôi màn "Lương thưởng & Phúc lợi" (tc-admin-api /metrics/hr/payroll) cho những THÁNG CHƯA CÓ
báo cáo `Bctienluong` (spec `hcns_luong_phongban` -> HCNS_LUONG_PB). Tháng nào có cả hai thì
màn ưu tiên Bctienluong — đó là báo cáo HCNS đã chốt, gồm cả các bảng thưởng chi riêng.

NGUỒN: `/data/9.HT_HCNS/1.1 TAI LIEU QUAN TRI/1.BAO CAO LUONG/BANG LUONG NAM <yyyy>/Bang luong
thang <mm>.<yyyy>[/...]` — trên 252 là `received_reports/HCNS/BANG LUONG NAM <yyyy> - Bang luong
thang <mm>.<yyyy>[ - ...]/`. Mỗi tháng một workbook gộp mọi công ty (sheet 'TH' + 'Bảng lương'
theo từng nhân viên + 'BHXH'). KỲ lấy từ TÊN THƯ MỤC, không từ tên file: tên file đổi mỗi tháng
('Bảng lương all T6.2026_SĐTCM', 'B.9.GR.HCNS.2026.7.BangluongAllT7.2026', và T9 là 'Báo cáo tổng
hợp chi phí lương và BHXH full theo SĐTC mới' — không mang kỳ).

VÌ SAO DERIVER, KHÔNG PHẢI SPEC JSON:
  · Các cột tổng hợp ('Lương CB'…'Thực lĩnh') nằm ở KHỐI CUỐI bảng (cột DS..EK), nhưng tên 'Thực
    lĩnh', 'Thuế', 'BH' cũng có ở phần chi tiết phía trước. Spec lấy cột đầu tiên trùng tên ->
    đọc nhầm. Ở đây lấy các cột SAU 'Lương CB' CUỐI CÙNG của dòng tiêu đề.
  · Dòng tiêu đề trôi: dòng 9 ở T4–T8, dòng 13/16 ở bản T9; có bản 3 dòng cùng chứa 'MÃ NV'.
    Nhận dòng có ĐỦ 'MÃ NV' + 'Lương CB' + 'Thực lĩnh'.
  · Phải CHỌN một file mỗi tháng: thư mục tháng có cả bản thường lẫn bản '_SĐTCM', và bản
    'Bảng lương all T9' là MẪU RỖNG (toàn 0) trong khi số T9 thật nằm ở file 'theo SĐTC mới'.

KHÔNG TIN SHEET 'TH': nó là SUMIFS theo khoá Công ty&Khối&Ban, và đo 17/10/2026 thì bản T9 ghi
'Tháng 08' với toàn số 0, bản T8 ra BHXH công ty ÂM. Cộng thẳng từ sheet nhân viên.
KHÔNG DÙNG SHEET 'BHXH': tên công ty viết khác sheet lương ('Khai thác' hoa), thiếu cả công ty
ở nhiều tháng, bản T9 để trống.

CHỈ TIÊU (đồng):
  thucLinh (amount2) = Tổng CP − giảm trừ − BH − công đoàn − thuế + hoàn thuế 5%. KHÔNG dùng cột
                       'Thực lĩnh' từng người: cột đó đã trừ thêm 'Giữ lương/lương âm', còn
                       Bctienluong thì chưa. Đo T6: khớp Bctienluong 7/7 công ty tới đồng; T8 khớp
                       'Bao cao tong hop luong' của HCNS 5/5 công ty kiểm được.
  thuong             = Thưởng thành tích + Thưởng KPI + Thưởng chuyên cần. KHÔNG gồm 'Thưởng
                       TN&Doanh số' — với tài xế Xanh VP / An Taxi đó là phần ăn chia doanh thu
                       (10–12 tỷ/tháng ở XVP), không phải thưởng.
  thuNhap (amount)   = Tổng CP − thuong, để thuNhap + thuong = Tổng CP (#27 Tổng quan cộng hai số
                       này; để thuNhap = Tổng CP là đếm thưởng hai lần).
  bhxh               = BH người lao động × 32/10,5 (21,5% công ty + 10,5% NLĐ theo luật). Khớp
                       'Tổng BHXH' của Bctienluong trong 1% ở cả 7 công ty T6.
  lvCuoiKy           = số người trên bảng lương KHÔNG có ngày nghỉ việc ≤ cuối tháng.

KHỐI: cột 'Khối ' / 'Phòng/Bộ phận' ở đầu bảng (cột 'Khối chức năng' của khối tổng hợp để trống
từng dòng). Theo sơ đồ tổ chức mới, Hưng Thịnh chỉ ghi 'Khối Hưng Thịnh' và dời khối thật sang
cột phòng ('Khối các Dự án Xây dựng', 'Khối kinh doanh xe tải'…) -> với các khối mang tên công ty
thì lấy phòng làm khối. Ánh xạ sang khối của bộ lọc nằm ở tc-admin-api app/data/khoi_map.json.

CÔNG TY: tên trong bảng lương đổi theo tháng ('Vinfast Thịnh Cường' T4–T7, 'Chi nhánh Vinfast -
Công ty Thịnh Cường' từ T8) -> quy về đúng tên pháp nhân dashboard (company_map.json `hcns_name`).

GHI: mỗi tháng đúng MỘT file. source_file = 'HCNS::<tên file>#<yyyy-mm>' (tên file 'Báo cáo tổng
hợp ... theo SĐTC mới' có thể lặp lại ở tháng sau). Trước khi ghi xoá dòng HCNS_LUONG_BL của kỳ
đó từ file KHÁC — đổi lựa chọn file không được để lại bản cũ cộng đôi.

Chạy: .venv/bin/python scripts/derive_hcns_bangluong.py [--thang 2026-07 ...] [--write]
"""
import argparse
import calendar
import collections
import datetime as dt
import glob
import json
import os
import re
import sys
import unicodedata

import openpyxl
import psycopg

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)
import spec_extract as SE  # noqa: E402  — dùng chung REPORTS_DIR/_ghi/_source_file_ghi/DB_URL

RT = "HCNS_LUONG_BL"
ROW_INDEX_BASE = 7887000
TY_LE_BHXH = 32 / 10.5

# Tên công ty trong bảng lương -> tên pháp nhân dashboard. Khoá là `_k()` của tên gốc.
_TEN_CONG_TY = {
    "vinfastthinhcuong": "Chi nhánh Vinfast - Công ty Thịnh Cường",
    "chinhanhvinfastcongtythinhcuong": "Chi nhánh Vinfast - Công ty Thịnh Cường",
    "vinfastuongbi": "Công ty Cổ phần Công Nghệ Vinfast Quảng Ninh",
    "congtycophancongnghevinfastquangninh": "Công ty Cổ phần Công Nghệ Vinfast Quảng Ninh",
    "congtycpcongnghevadichvuxanhvinhphuc": "Công ty Xanh Vĩnh Phúc",
    "congtyxanhvinhphuc": "Công ty Xanh Vĩnh Phúc",
    "congtycophanthinhcuong": "Công ty Cổ phần Thịnh Cường",
    "congtytnhhxuatnhapkhauvakhaithachungthinh": "Công ty TNHH Xuất nhập khẩu và khai thác Hưng Thịnh",
    "congtycophanglobalai": "Công ty Cổ phần Global AI",
    "congtycophanananhsgarden": "Công ty Cổ phần An An's Garden",
    "congtycophananansgarden": "Công ty Cổ phần An An's Garden",
}
# Khối mang tên CÔNG TY (sơ đồ tổ chức mới) -> khối thật nằm ở cột phòng.
_KHOI_LA_CONG_TY = {"khoihungthinh"}

_TIEN = ("luongcb", "luongythucchatluongcongviec", "tongphucap", "thuongthanhtich", "thuongkpi",
         "thuongtndoanhso", "thuongchuyencan", "cackhoanphucloi", "truythutruylinh",
         "cackhoangiamtru", "bh", "congdoan", "thue", "hoanthue5")
_CP = _TIEN[:9]
_THUONG = ("thuongthanhtich", "thuongkpi", "thuongchuyencan")
_KY_THU_MUC = re.compile(r"luong thang (\d{1,2})\.(\d{4})")


def _k(s) -> str:
    """Khoá so tên: bỏ dấu, bỏ mọi ký tự không phải chữ/số, thường hoá."""
    s = unicodedata.normalize("NFD", str(s or "").replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn" and c.isalnum()).lower()


def _khong_dau(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "").replace("đ", "d").replace("Đ", "D"))
    return " ".join("".join(c for c in s if unicodedata.category(c) != "Mn").lower().split())


def _so(v) -> float:
    if isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("\xa0", "").replace(",", "").strip() or 0)
    except ValueError:
        return 0.0


def _chu(v) -> str:
    s = str(v if v is not None else "").strip()
    return "" if s in ("0", "None", "#N/A") else s


def ky_cua_thu_muc(ten: str) -> str | None:
    m = _KY_THU_MUC.search(_khong_dau(ten))
    return f"{m.group(2)}-{int(m.group(1)):02d}" if m else None


def _sua_luc(path: str) -> str:
    """Ngày sửa ở NGUỒN (sidecar .json của máy gửi), lùi về mtime trên 252."""
    try:
        with open(os.path.splitext(path)[0] + ".json", encoding="utf-8") as f:
            v = json.load(f).get("modified_at")
        if v:
            return str(v)
    except (OSError, ValueError):
        pass
    return dt.datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d %H:%M:%S")


def doc_file(path: str, ky: str) -> dict | None:
    """Đọc 1 workbook 'all'. None = không phải bảng lương gộp (thiếu sheet/tiêu đề, < 5 công ty)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:                                    # noqa: BLE001 — file hỏng: bỏ qua, báo ở out
        return None
    try:
        if "th" not in {_k(t) for t in wb.sheetnames}:
            return None
        ws_l = [ws for ws in wb.worksheets if "bangluong" in _k(ws.title)]
        if not ws_l:
            return None
        ws = max(ws_l, key=lambda w: w.max_row or 0)
        nam, thang = int(ky[:4]), int(ky[5:7])
        cuoi = dt.date(nam, thang, calendar.monthrange(nam, thang)[1])
        cot = goc = None
        dong_td = None
        nhom = collections.defaultdict(collections.Counter)
        cty_la, n_nv = set(), 0
        # max_col: vài bản khai dimension 16.383 cột ảo — không chặn thì mỗi dòng đọc 16k ô.
        for i, r in enumerate(ws.iter_rows(max_col=200, values_only=True), 1):
            if cot is None:
                k = [_k(c) for c in r]
                if "manv" in k and "luongcb" in k and "thuclinh" in k:
                    cb = max(j for j, x in enumerate(k) if x == "luongcb")
                    cot = {}
                    for j in range(cb, len(k)):
                        if k[j] in _TIEN and k[j] not in cot:
                            cot[k[j]] = j
                    goc = {x: k.index(x) for x in ("congty", "khoi", "phongbophan", "manv",
                                                   "ngaynghiviec") if x in k}
                    dong_td = i
                    if set(_TIEN) - set(cot) or {"congty", "khoi", "manv"} - set(goc):
                        return {"loi": f"thiếu cột {sorted(set(_TIEN) - set(cot))} "
                                       f"{sorted({'congty', 'khoi', 'manv'} - set(goc))}"}
                elif i > 40:
                    return None
                continue
            if not _chu(r[goc["manv"]]):
                continue
            cty_goc = _chu(r[goc["congty"]])
            if not cty_goc:
                continue
            cty = _TEN_CONG_TY.get(_k(cty_goc), cty_goc)
            if _k(cty_goc) not in _TEN_CONG_TY:
                cty_la.add(cty_goc)
            khoi = _chu(r[goc["khoi"]])
            phong = _chu(r[goc["phongbophan"]]) if "phongbophan" in goc else ""
            if _k(khoi) in _KHOI_LA_CONG_TY and phong:
                khoi, phong = phong, ""
            v = {x: _so(r[j]) for x, j in cot.items()}
            nv = r[goc["ngaynghiviec"]] if "ngaynghiviec" in goc else None
            nghi = isinstance(nv, (dt.datetime, dt.date)) and (
                nv.date() if isinstance(nv, dt.datetime) else nv) <= cuoi
            c = nhom[(cty, khoi, phong)]
            c["soNguoi"] += 1
            c["lvCuoiKy"] += 0 if nghi else 1
            for x in _TIEN:
                c[x] += v[x]
            n_nv += 1
    finally:
        wb.close()
    if cot is None:
        return None
    if not n_nv:
        # Có khuôn bảng lương nhưng không một dòng nhân viên (bản 'Bảng lương all T9' đo 17/10).
        return {"loi": "mẫu rỗng — không có dòng nhân viên nào"}
    if len({k[0] for k in nhom}) < 5:
        return None                                     # bảng lương MỘT công ty, không phải 'all'
    recs = []
    for (cty, khoi, phong), c in sorted(nhom.items()):
        cp = sum(c[x] for x in _CP)
        thuong = sum(c[x] for x in _THUONG)
        tl = cp - c["cackhoangiamtru"] - c["bh"] - c["congdoan"] - c["thue"] + c["hoanthue5"]
        recs.append({
            "ngay": f"{ky}-01", "dim1": khoi or None, "dim2": phong or None, "dim3": cty,
            "amount": cp - thuong, "amount2": tl,
            "payload": {"unit": "dong", "thuong": thuong, "thuongTnDoanhSo": c["thuongtndoanhso"],
                        "tongCP": cp, "bhxh": c["bh"] * TY_LE_BHXH, "bhxhNld": c["bh"],
                        "thueTncn": c["thue"], "lvCuoiKy": c["lvCuoiKy"],
                        "soNguoiBangLuong": c["soNguoi"]},
        })
    return {"recs": recs, "so_nv": n_nv, "dong_tieu_de": dong_td,
            "thuc_linh": sum(r["amount2"] for r in recs), "cong_ty_la": sorted(cty_la)}


def quet(chi_thang: set | None = None) -> dict[str, list[str]]:
    """{kỳ: [file .xlsx/.xlsm trong các thư mục 'Bang luong thang MM.YYYY*']}."""
    goc = os.path.join(SE.REPORTS_DIR, "HCNS")
    out = collections.defaultdict(list)
    for d in sorted(glob.glob(os.path.join(goc, "BANG LUONG NAM *"))):
        ky = ky_cua_thu_muc(os.path.basename(d))
        if not ky or (chi_thang and ky not in chi_thang):
            continue
        for p in sorted(glob.glob(os.path.join(d, "*.xls[xm]"))):
            if not os.path.basename(p).startswith("~$"):
                out[ky].append(p)
    return out


def chon(ky: str, files: list[str]) -> tuple[str | None, dict | None, list]:
    """Chọn MỘT file cho kỳ: bảng lương gộp có số; nhiều bản thì bản sửa sau cùng ở nguồn,
    hoà thì ưu tiên bản theo sơ đồ tổ chức mới (SĐTC)."""
    ung, bo = [], []
    for p in files:
        d = doc_file(p, ky)
        if not d:
            continue
        if d.get("loi"):
            bo.append({"file": os.path.basename(p), "ly_do": d["loi"]})
        elif abs(d["thuc_linh"]) < 1:
            bo.append({"file": os.path.basename(p), "ly_do": "mẫu rỗng (thực lĩnh = 0)"})
        else:
            ung.append((_sua_luc(p), "sdtc" in _k(os.path.basename(p)), os.path.basename(p), p, d))
    if not ung:
        return None, None, bo
    ung.sort()
    for x in ung[:-1]:
        bo.append({"file": x[2], "ly_do": f"có bản khác chọn trước (sửa {x[0]})"})
    return ung[-1][3], ung[-1][4], bo


def ghi(ky: str, path: str, recs: list) -> dict:
    spec = {"report_type": RT, "row_index_base": ROW_INDEX_BASE, "source_file_hau_to": ky}
    sf = SE._source_file_ghi(spec, path)
    with psycopg.connect(SE.DB_URL) as conn:
        n = conn.execute("DELETE FROM raw_rows WHERE report_type=%s AND period_month=%s "
                         "AND source_file <> %s", (RT, ky, sf)).rowcount
        conn.commit()
    r = SE._ghi(spec, path, recs)
    if n:
        r["xoa_ban_file_khac"] = n
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--thang", nargs="*", help="chỉ các kỳ YYYY-MM")
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    ket_qua = []
    for ky, files in sorted(quet(set(a.thang) if a.thang else None).items()):
        path, d, bo = chon(ky, files)
        o = {"ky": ky, "file": os.path.basename(path) if path else None, "bo": bo}
        if d:
            tong = collections.defaultdict(collections.Counter)
            for r in d["recs"]:
                t = tong[r["dim3"]]
                t["thuNhap"] += r["amount"]
                t["thucLinh"] += r["amount2"]
                for x in ("thuong", "bhxh", "lvCuoiKy"):
                    t[x] += r["payload"][x]
            o.update({"dong": len(d["recs"]), "so_nv": d["so_nv"],
                      "cong_ty_la": d["cong_ty_la"],
                      "theo_cong_ty_ty": {c: {x: round(v / 1e9, 3) if x != "lvCuoiKy" else int(v)
                                              for x, v in t.items()} for c, t in tong.items()}})
            if a.write:
                o["ghi"] = ghi(ky, path, d["recs"])
        ket_qua.append(o)
    print(json.dumps(ket_qua, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
