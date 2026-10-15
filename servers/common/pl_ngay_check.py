# -*- coding: utf-8 -*-
"""Truy NGUỒN số P&L của MỘT NGÀY: DB đang có gì, file ngày đó có mấy bản, sheet KQKD từng bản ghi gì (15/10/2026).

Câu hỏi thật làm bot bí: "Tổng chi phí An KS 30/09 đang hiện 12.950.064 đ, chỉ nguồn lấy số". Dashboard chỉ lưu dòng gộp
(CHIPHI_D 'Chi phí chung' = CPBH + QLDN) nên muốn tách QLDN phải mở chính file. Ngày đó lại có HAI bản (bản KT sửa tay
`..xlsx` và bản tên chuẩn xuất lại sau) — `derive_hqkd_ngay._BAN_SUA_TAY` quyết định bản nào lên dashboard. Tool trả đủ:
dòng DB + từng bản file (ai sửa, lúc nào, đang được dùng hay không) + các dòng KQKD theo MÃ SỐ + ô khác nhau giữa các bản.
Chỉ đọc."""
import glob
import os
import re
import sys

from . import so_lieu_check as sc

# Mã số mẫu B02-DN cần cho câu hỏi chi phí/lợi nhuận.
_MA = {"01": "Doanh thu", "02": "Giảm trừ DT", "10": "DT thuần", "11": "Giá vốn", "20": "LN gộp", "21": "DT tài chính",
       "22": "CP tài chính", "25": "Chi phí bán hàng", "26": "Chi phí quản lý doanh nghiệp", "30": "LN thuần HĐKD",
       "31": "Thu nhập khác", "32": "Chi phí khác", "40": "LN khác", "50": "LNTT", "60": "LNST"}

# Cách người dùng gọi tên đơn vị -> tiền tố nguồn trong DB.
_BI_DANH_DV = {"ANKS": "ANKHACHSAN", "ANKHACHSAN": "ANKHACHSAN", "KHACHSAN": "ANKHACHSAN", "ANTAXI": "ANTAXI",
               "TRAMSAC": "TRAMSAC", "HUNGTHINH": "HUNGTHINH", "XETAI": "HUNGTHINH", "DUAN": "DUAN"}


def _ban_sua_tay():
    try:
        sys.path.insert(0, os.path.join(sc._AGENT_ROOT, "scripts"))
        from derive_hqkd_ngay import _BAN_SUA_TAY  # noqa: E402
        return dict(_BAN_SUA_TAY)
    except Exception:  # noqa: BLE001
        return {}


def _doc_kqkd(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = next((w for w in wb.worksheets if "kqkd" in w.title.lower()), None)
    meta = {"sua_boi": wb.properties.lastModifiedBy, "sua_luc": str(wb.properties.modified or ""),
            "tao_boi": wb.properties.creator}
    if ws is None:
        return meta, None, f"không có sheet KQKD (sheet: {', '.join(wb.sheetnames[:12])})"
    rows = [tuple(r) + (None,) * 12 for r in ws.iter_rows(values_only=True)]
    cot_ky = None
    for r in rows[:15]:
        for j, v in enumerate(r):
            if str(v or "").strip().lower() == "kỳ này":
                cot_ky = j
        if cot_ky is not None:
            break
    if cot_ky is None:
        return meta, None, "sheet KQKD không có cột 'Kỳ này'"
    dong = {}
    for i, r in enumerate(rows, 1):
        ma = str(r[1] or "").strip()
        if ma in _MA and ma not in dong:
            dong[ma] = {"dong": i, "chi_tieu": str(r[0] or "").strip()[:60], "ky_nay": r[cot_ky],
                        "cac_cot_so": [v for v in r[cot_ky:cot_ky + 7] if isinstance(v, (int, float))]}
    return meta, dong, None


def pl_ngay_nguon(don_vi: str, ngay: str, env: str = "prod", thu_muc: str = "baocaohqkdngay") -> dict:
    k = re.sub(r"[\s_.-]+", "", str(don_vi or "")).upper()
    don_vi, _ = sc._chuan_don_vi(env, _BI_DANH_DV.get(k, don_vi))
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", ngay or ""):
        raise ValueError("ngay phải dạng YYYY-MM-DD")
    db = sc._ro(env, "SELECT report_type, dim1, amount, source_file FROM raw_rows WHERE source_file LIKE %s AND ngay=%s "
                     "AND report_type IN ('PNLT_D','CHIPHI_D','HQKD_D','DTHU_D') AND COALESCE(amount,0)<>0 "
                     "ORDER BY 1, 2", (don_vi + "::%", ngay))
    ymd = ngay.replace("-", "")
    files = sorted(glob.glob(os.path.join(sc.RECEIVED_DIR, don_vi, thu_muc, f"*.D.{ymd}*.xls*")))
    sua = _ban_sua_tay()
    bi_thay = {k for k, v in sua.items() if os.path.exists(os.path.join(sc.RECEIVED_DIR, don_vi, thu_muc, v))}
    ban = []
    for p in files:
        ten = os.path.basename(p)
        meta, dong, loi = _doc_kqkd(p)
        ban.append({"file": ten, **meta, "ve_252_luc": sc.datetime.datetime.fromtimestamp(os.path.getmtime(p)).isoformat(" ", "seconds"),
                    "dashboard_dang_dung": ten not in bi_thay,
                    "ly_do": (f"bị thay bởi bản sửa tay {sua[ten]} (_BAN_SUA_TAY)" if ten in bi_thay
                              else "bản sửa tay được ghim" if ten in sua.values() else None),
                    "kqkd": dong, "loi": loi})
    khac = []
    if len(ban) > 1:
        goc = ban[0]["kqkd"] or {}
        for b in ban[1:]:
            for ma, d in (b["kqkd"] or {}).items():
                a = goc.get(ma)
                if a and a["cac_cot_so"] != d["cac_cot_so"]:
                    khac.append({"ma_so": ma, "chi_tieu": d["chi_tieu"], ban[0]["file"]: a["cac_cot_so"],
                                 b["file"]: d["cac_cot_so"]})
    out = {"don_vi": don_vi, "ngay": ngay, "env": env,
           "db": [{"report_type": r["report_type"], "chi_tieu": r["dim1"], "so_tien_dong": round((r["amount"] or 0) * 1e9),
                   "source_file": r["source_file"]} for r in db],
           "cac_ban_file": ban, "o_khac_nhau_giua_cac_ban": khac,
           "giai_thich": ("DB lưu số đã quy TỶ ĐỒNG; 'so_tien_dong' đã nhân lại ra đồng. CHIPHI_D 'Chi phí chung' = mã 25 (CPBH) "
                          "+ mã 26 (QLDN) cột 'Kỳ này' của sheet KQKD — dashboard KHÔNG lưu riêng QLDN, muốn tách thì đọc "
                          "kqkd['26'] của bản dashboard_dang_dung=true. Nhiều bản cùng ngày: chỉ bản dashboard_dang_dung lên số; "
                          "bản khác có số khác = kế toán đã sửa lại, cần người phụ trách quyết định đổi sang bản nào.")}
    if not files:
        out["canh_bao"] = f"không thấy file .D.{ymd} trong {don_vi}/{thu_muc}"
    return out
