# -*- coding: utf-8 -*-
"""SO KHỚP MAPPING ↔ LOGIC LÊN SỐ cho bot `ai_mapping_bot` (14/10/2026).

BÀI TOÁN: tài liệu mapping (Excel của kế toán/BA) mô tả "chỉ tiêu X lấy từ file nào, công thức gì"; logic thật nằm
ở builder của từng màn (`tc-admin-api/app/metrics/*`) + spec trích xuất (`extract_specs/*.json`). Hai bên không có mã
liên kết, và đã lệch nhau nhiều lần. Bot này đọc CẢ HAI và chỉ ra chỗ nghi lệch.

PHÂN VAI (luật số 1 của kiến trúc AI — xem KIENTRUC_AI_DULIEU_QUYETDINH.md): code ở đây ĐỌC và ĐỐI CHIẾU
TẤT ĐỊNH (tên file/sheet/mã, report_type, spec, số dòng có trong DB). Việc ĐỌC HIỂU công thức bằng chữ ↔ mã
nguồn là của LLM — nên mọi kết luận về công thức phải ghi rõ là "nghi ngờ/cần người xác nhận", không phải bằng chứng.

CHỈ ĐỌC: Excel + mã nguồn + DB (read_only). Không tool nào nhận SQL/đường dẫn tự do.

INPUT bot hiểu (người dùng nói tiếng Việt, tool nhận tham số có cấu trúc):
  · `man`      mã màn (vhkd1, at2, xdv0, ts3…)           — xem `map_danh_muc`
  · `tu_khoa`  tên chỉ tiêu / từ khoá trong mapping       — vd "tỷ lệ hoàn thành doanh thu"
  · `tai_lieu` mã tài liệu (QTTC, VHKD, XDV, DUAN, HTXT, ANTAXI, XANHVP, QLTS) — tuỳ chọn
  · `period`   YYYY-MM để xem dữ liệu có hay chưa         — tuỳ chọn
"""
import fnmatch
import glob
import json
import os
import re
import subprocess
import time
import unicodedata

import psycopg
from psycopg.rows import dict_row

_HERE = os.path.dirname(os.path.abspath(__file__))
_AGENT_ROOT = os.path.normpath(os.path.join(_HERE, "..", ".."))
INDEX = os.path.join(_AGENT_ROOT, "mapping_index.json")
SPEC_DIR = os.path.join(_AGENT_ROOT, "extract_specs")
DUMP_SCRIPT = os.path.join(_AGENT_ROOT, "scripts", "mapping_logic_dump.py")
API_DIR = os.environ.get("TC_API_DIR") or "/home/itadmin/apps/tc-console/tc-admin-api"
API_PY = os.path.join(API_DIR, ".venv", "bin", "python")
CACHE_DIR = os.environ.get("MAPPING_CACHE_DIR") or os.path.join(_AGENT_ROOT, "memory", "mapping_bot")
_PORT = {"prod": 5434, "test": 5435}

_cache_map = {}


def _kd(s):
    """Không dấu, thường, gọn khoảng trắng — để so khớp tiếng Việt."""
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.lower()).strip()


def _chi_muc():
    with open(INDEX, encoding="utf-8") as fh:
        return json.load(fh)


# ───────────────────── LOGIC (builder) ─────────────────────
def _logic():
    """Dump logic mọi màn; cache theo mtime của thư mục metrics để không chạy lại vô ích."""
    os.makedirs(CACHE_DIR, exist_ok=True)
    cf = os.path.join(CACHE_DIR, "logic.json")
    mt = max((os.path.getmtime(p) for p in glob.glob(os.path.join(API_DIR, "app", "metrics", "*.py"))), default=0)
    if os.path.exists(cf) and os.path.getmtime(cf) >= mt:
        with open(cf, encoding="utf-8") as fh:
            return json.load(fh)
    r = subprocess.run([API_PY, DUMP_SCRIPT], capture_output=True, text=True, timeout=120,
                       env={**os.environ, "TC_API_DIR": API_DIR})
    if r.returncode != 0:
        raise RuntimeError("không dựng được bảng logic các màn: " + (r.stderr or "")[-300:])
    with open(cf, "w", encoding="utf-8") as fh:
        fh.write(r.stdout)
    return json.loads(r.stdout)


def _man_hop_le(man):
    lg = _logic()["man"]
    if man not in lg:
        raise ValueError(f"không có màn '{man}'. Màn hiện có: {', '.join(sorted(lg))}")
    return lg[man]


# ───────────────────── TÀI LIỆU MAPPING ─────────────────────
_ROLES = (("nhom_ct", r"^nhom chi tieu"), ("man_hinh", r"^man hinh"), ("ma", r"ma chi tieu|^ma so|^ma$"),
          ("chi_tieu", r"^chi tieu( phan tich)?|truong tren giao dien|^chi tieu"),
          ("cong_thuc", r"cong thuc"), ("nguon", r"duong dan|nguon du lieu|mo ta cach lay|cach lay"),
          ("chieu", r"chieu"), ("ghi_chu", r"ghi chu|y nghia|luu y"))


def _doc_sheet(file, sheet):
    """Chuẩn hoá một sheet mapping -> list dòng {nhom, chi_tieu, cong_thuc, nguon, ma, chieu, ghi_chu, dong}.

    Tài liệu mapping KHÔNG đồng nhất cột: dò dòng tiêu đề (có ô 'Chỉ tiêu'/'Trường trên giao diện'), gán vai cột
    theo từ khoá tiêu đề. Dòng chỉ có 1 ô có chữ là TIÊU ĐỀ NHÓM (vd 'DASHBOARD mẹ: …') — gán làm `nhom`."""
    import openpyxl
    p = os.path.join(_chi_muc()["thu_muc_tai_lieu"], file)
    khoa = (p, sheet, os.path.getmtime(p))
    if khoa in _cache_map:
        return _cache_map[khoa]
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ten_that = next((n for n in wb.sheetnames if _kd(n) == _kd(sheet)), None)   # tên sheet hay lẫn dấu cách/ký tự lạ
    if ten_that is None:
        raise ValueError(f"tài liệu '{file}' không có sheet '{sheet}'")
    rows = [list(r) for r in wb[ten_that].iter_rows(values_only=True)]
    hd, cot = None, {}
    for i, r in enumerate(rows[:10]):
        cells = [_kd(c) for c in r]
        if any(re.match(r"^chi tieu|truong tren giao dien", c) for c in cells if c):
            hd = i
            for j, c in enumerate(cells):
                for vai, rx in _ROLES:
                    if vai not in cot and c and re.search(rx, c):
                        cot[vai] = j
            break
    out, nhom = [], None
    if hd is None:
        _cache_map[khoa] = out
        return out
    for i, r in enumerate(rows[hd + 1:], start=hd + 2):
        co = [(j, c) for j, c in enumerate(r) if c is not None and str(c).strip() not in ("", "\xa0")]
        if not co:
            continue
        if len(co) == 1 and co[0][0] == 0 and not str(co[0][1]).strip().isdigit():
            nhom = str(co[0][1]).strip()[:160]
            continue

        def v(vai):
            j = cot.get(vai)
            return None if j is None or j >= len(r) or r[j] is None else str(r[j]).strip()
        ct = v("chi_tieu") or v("ma")
        if not ct or ct == "—":
            continue
        out.append({"dong": i, "nhom": nhom, "nhom_chi_tieu": v("nhom_ct"), "man_hinh": v("man_hinh"), "ma": v("ma"), "chi_tieu": ct,
                    "cong_thuc": v("cong_thuc"), "nguon": v("nguon"), "chieu": v("chieu"), "ghi_chu": v("ghi_chu")})
    _cache_map[khoa] = out
    return out


def _tai_lieu_cua_man(man):
    ci = _chi_muc()
    return [(k, d) for k, d in ci["tai_lieu"].items() if man in d["man"]]


# ───────────────────── SPEC + DB ─────────────────────
def _specs():
    out = []
    for f in sorted(glob.glob(os.path.join(SPEC_DIR, "*.json"))):
        try:
            with open(f, encoding="utf-8") as fh:
                s = json.load(fh)
        except (OSError, ValueError):
            continue
        ng = s.get("nguon") or {}
        out.append({"id": s.get("id"), "report_type": s.get("report_type"), "khoi": s.get("khoi"),
                    "folder": ng.get("folder"), "file_glob": ng.get("file_glob")})
    return out


def _db_dem(env, report_types, period=None):
    """{report_type: (số dòng, kỳ mới nhất có dữ liệu)}."""
    if not report_types:
        return {}
    from urllib.parse import unquote, urlparse
    dsn = os.environ.get("DATABASE_URL") or ""
    mk = unquote(urlparse(dsn).password or "") if dsn else ""
    with psycopg.connect(host="localhost", port=_PORT[env], dbname="tc_dashboard", user="tc",
                         password=mk or "tc_$production", autocommit=True, row_factory=dict_row,
                         prepare_threshold=None, options="-c default_transaction_read_only=on "
                                                         "-c statement_timeout=30000") as c:
        sql = ("SELECT report_type, COUNT(*) AS n, MAX(period_month) AS ky FROM raw_rows "
               "WHERE report_type = ANY(%s)" + (" AND period_month=%s" if period else "") + " GROUP BY report_type")
        rows = c.execute(sql, (sorted(report_types), period) if period else (sorted(report_types),)).fetchall()
    return {r["report_type"]: (r["n"], r["ky"]) for r in rows}


_RE_FILE = re.compile(r"[A-Za-z0-9][\w.\-]*\.(?:xlsx|xlsm|xlsb|xls)\b", re.IGNORECASE)


def _spec_cho_file(ten, specs):
    """Spec nào có file_glob khớp tên file (so thô: glob → fnmatch, không phân biệt hoa thường)."""
    return [s for s in specs if s["file_glob"] and fnmatch.fnmatch(ten.lower(), str(s["file_glob"]).lower())]


# ───────────────────── 1. DANH MỤC / ĐỘ PHỦ ─────────────────────
def danh_muc(env="prod", period=None):
    """Độ phủ mapping của từng màn: có tài liệu nào, builder đọc report_type nào, spec nào nạp, có dữ liệu chưa."""
    lg = _logic()["man"]
    specs = _specs()
    rt_spec = {}
    for s in specs:
        rt_spec.setdefault(s["report_type"], []).append(s["id"])
    tat = {rt for m in lg.values() for rt in m["report_types"]}
    dem = _db_dem(env, tat, period)
    ci = _chi_muc()
    man_ra = []
    for man, m in sorted(lg.items()):
        tl = [k for k, _ in _tai_lieu_cua_man(man)]
        rts = m["report_types"]
        man_ra.append({
            "man": man, "builder": m.get("builder"), "file_logic": f"{m.get('file')}:{m.get('dong_bat_dau')}",
            "tai_lieu_mapping": tl or None, "chi_co_tai_lieu_md": ci["chi_co_tai_lieu_md"].get(man),
            "so_report_type": len(rts),
            "rt_khong_co_spec_nap": [r for r in rts if r not in rt_spec and not r.endswith("_D") and r not in dem],
            "rt_khong_co_du_lieu": [r for r in rts if r not in dem],
            "danh_gia": ("CÓ tài liệu mapping" if tl else
                         "CHỈ có tài liệu .md, không có bảng mapping chuẩn" if man in ci["chi_co_tai_lieu_md"]
                         else "KHÔNG có tài liệu mapping")})
    thieu = [x["man"] for x in man_ra if x["danh_gia"].startswith("KHÔNG")]
    return {"env": env, "period": period, "so_man": len(man_ra), "man_khong_co_tai_lieu_mapping": thieu,
            "man_chi_co_md": sorted(ci["chi_co_tai_lieu_md"]),
            "tai_lieu": {k: {"ten": d["ten"], "file": d["file"], "so_man": len(d["man"])} for k, d in ci["tai_lieu"].items()},
            "man": man_ra,
            "luu_y": "Độ phủ này là 'có tài liệu hay không', KHÔNG phải 'mapping đúng hay không' — dùng map_doi_chieu."}


# ───────────────────── 2. ĐỌC MAPPING ─────────────────────
def tim_chi_tieu(tu_khoa, man=None, tai_lieu=None, toi_da=8):
    """Tìm dòng mapping theo từ khoá (không dấu). Giới hạn theo màn hoặc tài liệu nếu nêu."""
    ci = _chi_muc()
    kw = _kd(tu_khoa)
    if len(kw) < 3:
        raise ValueError("từ khoá quá ngắn")
    if man:
        _man_hop_le(man)
        dsl = [(k, ci["tai_lieu"][k]) for k, _ in _tai_lieu_cua_man(man)]
    elif tai_lieu:
        if tai_lieu not in ci["tai_lieu"]:
            raise ValueError(f"tài liệu '{tai_lieu}' không có. Có: {', '.join(ci['tai_lieu'])}")
        dsl = [(tai_lieu, ci["tai_lieu"][tai_lieu])]
    else:
        dsl = list(ci["tai_lieu"].items())
    ra = []
    for k, d in dsl:
        for sh in d["sheet"]:
            try:
                rows = _doc_sheet(d["file"], sh)
            except (OSError, ValueError):
                continue
            for r in rows:
                hay = _kd(" ".join(str(r.get(f) or "") for f in ("nhom", "nhom_chi_tieu", "chi_tieu", "ma", "man_hinh", "cong_thuc")))
                if kw in hay:
                    ra.append({"tai_lieu": k, "file": d["file"], "sheet": sh, **r})
    ra.sort(key=lambda r: 0 if kw in _kd(r["chi_tieu"]) else 1)
    return {"tu_khoa": tu_khoa, "so_khop": len(ra), "ket_qua": ra[:toi_da],
            "cat_bot": max(0, len(ra) - toi_da),
            "luu_y": "Mỗi tài liệu có nhiều bản/sheet (vd 'Mapping Antaxi 25.08' vs '13.08'); đã chọn sheet mới nhất đã khai."}


# ───────────────────── 3. ĐỌC LOGIC ─────────────────────
def logic_man(man, so_dong_ma=70):
    m = _man_hop_le(man)
    return {"man": man, "builder": m.get("builder"), "module": m.get("module"),
            "file": m.get("file"), "dong_bat_dau": m.get("dong_bat_dau"), "so_dong_ham": m.get("so_dong_ham"),
            "report_type_doc": m["report_types"], "nguon_ngay": m.get("rt_ngay"),
            "mo_ta_trong_code": m.get("doc"),
            "ma_nguon_dau_ham": "".join((m.get("ma_nguon") or "").splitlines(True)[:max(10, min(so_dong_ma, 160))]),
            "luu_y": "report_type rút bằng AST (đủ để biết màn ĐỌC nguồn nào; không phải công thức đầy đủ — "
                     "công thức chi tiết nằm trong thân hàm và các hàm phụ)."}


# ───────────────────── 4. ĐỐI CHIẾU ─────────────────────
def doi_chieu(man, tu_khoa, env="prod", period=None):
    """Một chỉ tiêu: dòng mapping ↔ logic ↔ spec ↔ dữ liệu. Phần tất định ghi cờ; phần công thức để người/LLM đọc."""
    m = _man_hop_le(man)
    tk = tim_chi_tieu(tu_khoa, man=man, toi_da=4)
    specs = _specs()
    rts_builder = set(m["report_types"])
    phan_tich = []
    for r in tk["ket_qua"]:
        ten_file = sorted(set(_RE_FILE.findall(r.get("nguon") or "")))
        spec_khop = {t: [s["id"] for s in _spec_cho_file(t, specs)] for t in ten_file}
        rt_tu_spec = {s["report_type"] for t in ten_file for s in _spec_cho_file(t, specs)}
        dem = _db_dem(env, rt_tu_spec | rts_builder, period) if (rt_tu_spec or rts_builder) else {}
        co = []
        if not ten_file:
            co.append("MAPPING_KHONG_NEU_TEN_FILE: mô tả nguồn không nêu tên file cụ thể — không đối chiếu tự động được")
        for t, ids in spec_khop.items():
            if not ids:
                co.append(f"KHONG_CO_SPEC: mapping nêu file '{t}' nhưng không spec nào khớp file_glob (nguồn có thể "
                          f"chưa được nạp tự động, hoặc nạp bằng deriver riêng)")
        if rt_tu_spec and not (rt_tu_spec & rts_builder):
            co.append(f"BUILDER_KHONG_DOC_NGUON_NAY: spec nạp vào {sorted(rt_tu_spec)} nhưng builder '{m.get('builder')}' "
                      f"không đọc report_type nào trong đó")
        thieu_dl = sorted(rt for rt in (rt_tu_spec & rts_builder) if rt not in dem)
        if thieu_dl:
            co.append(f"NGUON_CHUA_CO_DU_LIEU: {thieu_dl} chưa có dòng nào trong DB {env}"
                      + (f" kỳ {period}" if period else ""))
        phan_tich.append({"dong_mapping": {k: r[k] for k in ("tai_lieu", "sheet", "dong", "nhom", "nhom_chi_tieu", "chi_tieu", "ma",
                                                              "cong_thuc", "nguon", "chieu", "ghi_chu")},
                          "file_nguon_trong_mapping": ten_file, "spec_khop_file": spec_khop,
                          "report_type_tu_spec": sorted(rt_tu_spec),
                          "du_lieu_trong_db": {k: {"so_dong": v[0], "ky_moi_nhat": v[1]} for k, v in dem.items()},
                          "co_nghi_van": co})
    kw = [w for w in _kd(tu_khoa).split() if len(w) >= 4]
    dong_ma = []
    for i, ln in enumerate((m.get("ma_nguon") or "").splitlines(), start=m.get("dong_bat_dau") or 1):
        if kw and any(w in _kd(ln) for w in kw):
            dong_ma.append(f"{m.get('file')}:{i}: {ln.strip()[:160]}")
    return {"man": man, "tu_khoa": tu_khoa, "env": env, "period": period,
            "so_dong_mapping_khop": tk["so_khop"], "doi_chieu": phan_tich,
            "logic": {"builder": m.get("builder"), "file": f"{m.get('file')}:{m.get('dong_bat_dau')}",
                      "report_type_doc": sorted(rts_builder), "dong_ma_chua_tu_khoa": dong_ma[:20]},
            "can_nguoi_hoac_llm_doc": "So công thức bằng chữ (cong_thuc) với thân hàm builder: dùng map_logic_man để "
                                      "lấy mã nguồn. Kết luận về công thức chỉ là NGHI NGỜ, không phải bằng chứng."}


# ───────────────────── 5. TRUY VẾT MỘT CHỈ TIÊU ─────────────────────
_RE_MA = re.compile(r"\b(CT\d{2})\b")


def _file_that_trong_db(env, report_types, period=None, moi_rt=6):
    """{report_type: [{file, so_dong}]} — tên file nguồn THẬT đang nuôi report_type đó (bỏ tiền tố 'ĐƠN VỊ::')."""
    if not report_types:
        return {}
    from urllib.parse import unquote, urlparse
    dsn = os.environ.get("DATABASE_URL") or ""
    mk = unquote(urlparse(dsn).password or "") if dsn else ""
    with psycopg.connect(host="localhost", port=_PORT[env], dbname="tc_dashboard", user="tc",
                         password=mk or "tc_$production", autocommit=True, row_factory=dict_row,
                         prepare_threshold=None, options="-c default_transaction_read_only=on "
                                                         "-c statement_timeout=30000") as c:
        sql = ("SELECT report_type, source_file, COUNT(*) AS n FROM raw_rows WHERE report_type = ANY(%s)"
               + (" AND period_month=%s" if period else "") + " GROUP BY 1,2 ORDER BY 1, 3 DESC")
        rows = c.execute(sql, (sorted(report_types), period) if period else (sorted(report_types),)).fetchall()
    ra = {}
    for r in rows:
        lst = ra.setdefault(r["report_type"], [])
        if len(lst) < moi_rt:
            lst.append({"file": (r["source_file"] or "").split("::", 1)[-1], "so_dong": r["n"]})
    return ra


def truy_vet(man, tu_khoa, env="prod", period=None, ngu_canh=3):
    """Truy vết MỘT chỉ tiêu từ đầu đến cuối: dòng mapping (trích nguyên văn) → đoạn mã dashboard nhắc chỉ tiêu →
    mã chỉ tiêu/dữ liệu code dùng → file nguồn thật trong DB. Trả cả danh sách việc CHƯA kiểm được."""
    m = _man_hop_le(man)
    tk = tim_chi_tieu(tu_khoa, man=man, toi_da=3)
    # 1. mapping: trích nguyên văn
    trich = []
    cum = {_kd(tu_khoa)}
    for r in tk["ket_qua"]:
        trich.append({"tai_lieu": r["tai_lieu"], "file_mapping": r["file"], "sheet": r["sheet"], "dong": r["dong"],
                      "nguyen_van": {k: r.get(k) for k in ("nhom", "nhom_chi_tieu", "chi_tieu", "ma", "cong_thuc",
                                                           "nguon", "chieu", "ghi_chu") if r.get(k)}})
        for k in ("chi_tieu", "nhom_chi_tieu"):
            if r.get(k) and len(_kd(r[k])) >= 6:
                cum.add(_kd(r[k]))
    # 2. mã nguồn: cả module chứa builder
    duong = os.path.join(API_DIR, m.get("file") or "")
    try:
        with open(duong, encoding="utf-8") as fh:
            dong = fh.read().splitlines()
    except OSError:
        dong = []
    trung, ma_tim = [], []
    for i, ln in enumerate(dong):
        k = _kd(ln)
        if any(c in k for c in cum):
            trung.append(i)
            ma_tim += _RE_MA.findall(ln)
    ma_tim = list(dict.fromkeys(ma_tim))[:4]
    for i, ln in enumerate(dong):                      # mở rộng: dòng nhắc đúng mã chỉ tiêu (CT33…)
        if len(trung) >= 12:
            break
        if i not in trung and any(re.search(rf"\b{re.escape(x)}\b", ln) for x in ma_tim):
            trung.append(i)
    doan, da = [], set()
    for i in sorted(trung)[:12]:
        a, b = max(0, i - ngu_canh), min(len(dong), i + ngu_canh + 1)
        if all(j in da for j in range(a, b)):
            continue
        da.update(range(a, b))
        doan.append({"dong": f"{m.get('file')}:{i + 1}",
                     "ma": "\n".join(f"{j + 1}: {dong[j][:170]}" for j in range(a, b))})
    van_ban = "\n".join(d["ma"] for d in doan)
    # 3. file nguồn thật
    rt_that = [rt for rt in m["report_types"] if not re.fullmatch(r"CT\d{2}", rt)]   # CT33.. là mã chỉ tiêu, không phải report_type
    rt_nhac = sorted(rt for rt in rt_that if re.search(rf"[\"'`]{re.escape(rt)}[\"'`]", van_ban))
    rt_dung = rt_nhac + [rt for rt in rt_that if rt not in rt_nhac]
    ftd = _file_that_trong_db(env, rt_dung, period, moi_rt=3)
    # 4. cờ & những gì CHƯA kiểm
    co = []
    ten_map = sorted({t for x in trich for t in _RE_FILE.findall(x["nguyen_van"].get("nguon") or "")})
    if not trich:
        co.append("KHONG_CO_DONG_MAPPING: không tìm thấy chỉ tiêu này trong tài liệu mapping của màn")
    elif not ten_map:
        co.append("MAPPING_KHONG_NEU_TEN_FILE: dòng mapping không nói file nguồn — không thể xác nhận dashboard lấy đúng file mapping chỉ")
    if not doan:
        co.append("CODE_KHONG_NHAC_CHI_TIEU: không dòng mã nào trong module nhắc tên chỉ tiêu/mã — thử từ khoá khác")
    for t in ten_map:
        if not any(t.lower() in (f["file"] or "").lower() or fnmatch.fnmatch(f["file"].lower(), "*" + t.lower().split(".")[-2] + "*")
                   for lst in ftd.values() for f in lst):
            co.append(f"FILE_MAPPING_KHONG_THAY_TRONG_DB: mapping nêu '{t}' nhưng DB {env} không có source_file nào tương ứng "
                      f"ở các report_type của chỉ tiêu")
    if rt_nhac and any(rt not in ftd for rt in rt_nhac):
        co.append(f"NGUON_CHUA_CO_DU_LIEU: {sorted(set(rt_nhac) - set(ftd))} được nhắc trong mã nhưng không có dòng nào"
                  + (f" kỳ {period}" if period else ""))
    return {"man": man, "tu_khoa": tu_khoa, "env": env, "period": period,
            "1_mapping_ghi": trich, "so_dong_mapping_khop": tk["so_khop"],
            "2_code_dashboard": {"builder": m.get("builder"), "file": f"{m.get('file')}:{m.get('dong_bat_dau')}",
                                 "ma_chi_tieu_trong_code": ma_tim, "doan_ma": doan},
            "3_file_nguon_that": {"rt_duoc_nhac_trong_doan_ma": rt_nhac,
                                  "rt_builder_con_lai": [rt for rt in rt_that if rt not in rt_nhac],
                                  "file_trong_db": ftd},
            "co_nghi_van": co,
            "CHUA_KIEM": ["công thức mapping (chữ) có cùng nghĩa với công thức trong đoạn mã hay không — chỉ đọc, chưa chạy số",
                          "con số đang hiển thị trên dashboard có bằng số tính tay từ file hay không (bot không gọi API hiển thị)"]}
