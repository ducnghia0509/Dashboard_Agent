# -*- coding: utf-8 -*-
"""ĐỌC EXCEL CHỈ ĐỌC cho bot `tim_loi` (15/10/2026, hợp đồng đợt 2+3 mục C).

VÌ SAO: bot tìm lỗi chạy bằng Claude Code thì "tự tính số đúng từ file" bằng openpyxl tự do. Trong chat (OpenClaw) bot
KHÔNG có exec — nên phép cộng có lọc phải nằm ở đây, tất định: `tong_hop_excel` là đường DUY NHẤT để bot ra một "số đúng
theo mapping" từ file nguồn. LLM chỉ chọn cột/điều kiện, không tự cộng.

AN TOÀN:
  · Đường file chỉ được nằm trong `received_reports` hoặc `Tài liệu` (GOC_CHO_PHEP). Cấm thành phần `..`, cấm symlink
    (realpath phải trùng đường đã chuẩn hoá — symlink trỏ ra ngoài gốc hay trỏ trong gốc đều bị từ chối, vì không có
    lý do chính đáng để đọc qua symlink), chỉ nhận đuôi .xlsx/.xlsm/.xlsb/.xls.
  · Chỉ ĐỌC (openpyxl read_only + data_only / pyxlsb / xlrd). Không ghi file, không exec.

QUY ƯỚC SO KHỚP (ghi rõ để người đọc kết quả biết bot đã lọc thế nào):
  · Chuỗi: chuẩn hoá Unicode NFC, bỏ dấu cách thừa (đầu/cuối + gộp khoảng trắng giữa), KHÔNG phân biệt hoa thường
    (casefold). GIỮ dấu tiếng Việt ('Đã nhận' = 'đã nhận ' nhưng ≠ 'Da nhan'). Áp cho cả tên cột (header) lẫn giá trị.
  · Ngày: nhận datetime/date của Excel, số serial Excel (1..2958465, gốc 1899-12-30 — đúng cách Excel/pyxlsb/xlrd trả),
    chuỗi gõ tay dd/mm/yyyy, d/m/yy, dd-mm-yyyy, dd.mm.yyyy, yyyy-mm-dd (kèm giờ phía sau cũng được). Chuỗi kiểu
    dd/mm luôn đọc NGÀY TRƯỚC THÁNG (quy ước Việt Nam). Ô không đọc được ngày bị coi là KHÔNG khớp điều kiện ngày và
    được đếm vào `o_ngay_khong_doc_duoc` (kèm vài mẫu) — không âm thầm bỏ.
  · Số: cộng bằng `Decimal(str(x))` — không làm tròn, không cộng dồn sai số float. Ô chữ dạng số ('1.720.000',
    '1,720,000.5', '12,5') được đọc — chuỗi có nhóm '.ddd' đầu khác 0 ('1.000', '12.500') đọc theo kiểu VN (chấm = phân cách
    nghìn), còn '1.5', '0.100' là thập phân; ô chữ không phải số / ô lỗi ('#REF!') KHÔNG cộng và được đếm vào `o_khong_phai_so`.

CACHE: workbook theo (đường dẫn, mtime, cỡ file, sheet) — LRU 8 sheet. File bị ghi đè (đổi mtime) là tự đọc lại.
"""
import datetime
import os
import re
import threading
import unicodedata
from collections import OrderedDict
from decimal import Decimal, InvalidOperation

ROOT = os.environ.get("TC_ROOT") or "/home/itadmin/AI_Dashboard_QT"
RECEIVED_DIR = os.environ.get("RECEIVED_DIR") or os.path.join(ROOT, "Connect_VPS", "received_reports")
TAI_LIEU_DIR = os.environ.get("TAI_LIEU_DIR") or os.path.join(ROOT, "Tài liệu")
# Thứ tự = thứ tự thử khi người gọi đưa đường TƯƠNG ĐỐI. Test đơn vị thay biến này bằng thư mục tạm.
GOC_CHO_PHEP = [RECEIVED_DIR, TAI_LIEU_DIR]
DUOI_EXCEL = (".xlsx", ".xlsm", ".xlsb", ".xls")

GIOI_HAN_DONG_TRA = 500
GIOI_HAN_DONG_GOP = 40
_LRU_MAX = 8
_lru = OrderedDict()
_lru_lock = threading.Lock()


# ───────────────────────────── chuẩn hoá ─────────────────────────────
def chuan(s):
    """NFC + gộp khoảng trắng + casefold. GIỮ dấu. Dùng so cột/giá trị chữ."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFC", str(s)).replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip().casefold()


def _kd(s):
    """Không dấu (chỉ dùng tìm sheet/từ khoá gần đúng — KHÔNG dùng cho điều kiện lọc)."""
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.lower()).strip()


def json_gia_tri(v):
    """Giá trị ô -> JSON được (ngày -> ISO, Decimal -> số)."""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.isoformat()
    if isinstance(v, Decimal):
        return int(v) if v == v.to_integral_value() else float(v)
    return v


# ───────────────────────────── đường dẫn an toàn ─────────────────────────────
def duong_an_toan(file, duoi=DUOI_EXCEL):
    """Đường người gọi đưa -> đường tuyệt đối đã kiểm. Tương đối thì thử lần lượt từng gốc cho phép.

    Từ chối (ValueError): rỗng, có byte NUL, có thành phần '..', nằm ngoài gốc, có symlink ở bất kỳ thành phần nào
    dưới gốc, sai đuôi, không phải file thường."""
    s = str(file or "").strip()
    if not s or "\x00" in s:
        raise ValueError("thiếu đường file")
    if ".." in re.split(r"[\\/]+", s):
        raise ValueError("đường file không được chứa '..'")
    if duoi and not s.lower().endswith(tuple(duoi)):
        raise ValueError(f"chỉ đọc file đuôi {', '.join(duoi)}")
    ung_vien = [os.path.normpath(s)] if os.path.isabs(s) else [os.path.normpath(os.path.join(g, s)) for g in GOC_CHO_PHEP]
    loi = "file phải nằm trong received_reports hoặc Tài liệu"
    for p in ung_vien:
        g = next((os.path.normpath(g) for g in GOC_CHO_PHEP if p.startswith(os.path.normpath(g) + os.sep)), None)
        if g is None:
            continue
        if os.path.realpath(p) != os.path.join(os.path.realpath(g), os.path.relpath(p, g)):
            raise ValueError("đường file đi qua symlink — bị chặn")
        if os.path.isfile(p):
            return p
        loi = f"không có file '{s}'"
    raise ValueError(loi)


# ───────────────────────────── đọc workbook ─────────────────────────────
def _ten_sheet(p):
    low = p.lower()
    if low.endswith(".xlsb"):
        from pyxlsb import open_workbook
        with open_workbook(p) as wb:
            return list(wb.sheets)
    if low.endswith(".xls"):
        import xlrd
        return xlrd.open_workbook(p, on_demand=True).sheet_names()
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def ds_sheet(p):
    return _ten_sheet(p)


def _chon_sheet(p, sheet):
    ds = _ten_sheet(p)
    if sheet is None or str(sheet).strip() == "":
        raise ValueError(f"thiếu sheet. File có: {ds}")
    for n in ds:
        if n == sheet:
            return n
    for n in ds:
        if chuan(n) == chuan(sheet):
            return n
    for n in ds:
        if _kd(n) == _kd(sheet):
            return n
    raise ValueError(f"không có sheet '{sheet}'. File có: {ds}")


def _doc_tho(p, ten):
    low = p.lower()
    rows = []
    if low.endswith(".xlsb"):
        from pyxlsb import open_workbook
        with open_workbook(p) as wb:
            with wb.get_sheet(ten) as ws:
                for r in ws.rows():
                    out = []
                    for c in r:
                        while len(out) < c.c:
                            out.append(None)
                        out.append(c.v)
                    rows.append(out)
    elif low.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(p, on_demand=True)
        ws = wb.sheet_by_name(ten)
        for i in range(ws.nrows):
            out = []
            for c in ws.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    try:
                        out.append(xlrd.xldate.xldate_as_datetime(c.value, wb.datemode))
                    except Exception:  # noqa: BLE001
                        out.append(c.value)
                elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    out.append(None)
                elif c.ctype == xlrd.XL_CELL_ERROR:
                    out.append("#LỖI")
                else:
                    out.append(c.value)
            rows.append(out)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            for r in wb[ten].iter_rows(values_only=True):
                rows.append(list(r))
        finally:
            wb.close()
    for r in rows:                                   # cắt cột ảo cuối dòng (sheet hay có max_column hàng nghìn)
        while r and (r[-1] is None or (isinstance(r[-1], str) and r[-1].strip() == "")):
            r.pop()
    return rows


def doc_sheet(file, sheet, bo_cache=False):
    """-> (đường đã kiểm, tên sheet thật, list dòng [list ô]); dòng i (0-based) = dòng Excel i+1."""
    p = duong_an_toan(file)
    ten = _chon_sheet(p, sheet)
    st = os.stat(p)
    khoa = (p, st.st_mtime_ns, st.st_size, ten)
    with _lru_lock:
        if not bo_cache and khoa in _lru:
            _lru.move_to_end(khoa)
            return p, ten, _lru[khoa]
    rows = _doc_tho(p, ten)
    with _lru_lock:
        _lru[khoa] = rows
        _lru.move_to_end(khoa)
        while len(_lru) > _LRU_MAX:
            _lru.popitem(last=False)
    return p, ten, rows


# ───────────────────────────── cột / ô ─────────────────────────────
def chu_cot(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def so_cot(chu):
    n = 0
    for ch in chu.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


def header_map(rows, header_dong):
    """{tên chuẩn hoá: chỉ số cột} + danh sách tiêu đề nguyên văn. Tiêu đề trùng: bản sau mang hậu tố ' (2)'."""
    if not isinstance(header_dong, int) or header_dong < 1 or header_dong > len(rows):
        raise ValueError(f"header_dong phải là số dòng Excel 1..{len(rows)}")
    hd = rows[header_dong - 1]
    m, ds = {}, []
    for j, h in enumerate(hd):
        if h is None or str(h).strip() == "":
            continue
        ten = str(h).strip()
        k = chuan(ten)
        if k in m:
            n = 2
            while f"{k} ({n})" in m:
                n += 1
            k, ten = f"{k} ({n})", f"{ten} ({n})"
        m[k] = j
        ds.append({"cot": chu_cot(j), "tieu_de": ten})
    return m, ds


def tim_cot(hmap, cot):
    """Tên cột (so theo `chuan`) hoặc chữ cột Excel ('L'). Tên ưu tiên hơn chữ."""
    k = chuan(cot)
    if k in hmap:
        return hmap[k]
    if re.fullmatch(r"[A-Za-z]{1,3}", str(cot or "").strip()):
        return so_cot(str(cot).strip())
    raise ValueError(f"không có cột '{cot}'. Tiêu đề có: {sorted(hmap)}")


_SERIAL_GOC = datetime.datetime(1899, 12, 30)


def doc_ngay(v):
    """Ô -> datetime.date hoặc None. Xem quy ước ở docstring module."""
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        x = float(v)
        if 1 <= x < 2958466:
            return (_SERIAL_GOC + datetime.timedelta(days=int(x))).date()
        return None
    s = str(v).strip()
    if not s:
        return None
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ T].*)?$", s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})(?:\s.*)?$", s)
        if not m:
            if re.fullmatch(r"\d{5}(\.\d+)?", s):                   # serial gõ thành chữ
                return doc_ngay(float(s))
            return None
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
    try:
        return datetime.date(y, mo, d)
    except ValueError:
        return None


def doc_so(v):
    """Ô -> Decimal hoặc None (không phải số)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, Decimal):
        return v
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return None
        return Decimal(repr(v))
    if isinstance(v, (datetime.date, datetime.datetime, datetime.time)):
        return None
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    if not s:
        return None
    am = s.startswith("(") and s.endswith(")")
    if am:
        s = s[1:-1]
    try:
        if re.fullmatch(r"[-+]?[1-9]\d{0,2}(\.\d{3})+(,\d+)?", s):     # 1.720.000,5 / '1.000' (kiểu VN: chấm = nghìn)
            d = Decimal(s.replace(".", "").replace(",", "."))
        elif re.fullmatch(r"[-+]?\d+(\.\d+)?", s):
            d = Decimal(s)
        elif re.fullmatch(r"[-+]?\d{1,3}(,\d{3})+(\.\d+)?", s):         # 1,720,000.5
            d = Decimal(s.replace(",", ""))
        elif re.fullmatch(r"[-+]?\d+,\d+", s):                          # 12,5
            d = Decimal(s.replace(",", "."))
        else:
            return None
    except InvalidOperation:
        return None
    return -d if am else d


def _so_ra(d):
    return int(d) if d == d.to_integral_value() else float(d)


# ───────────────────────────── điều kiện lọc ─────────────────────────────
TOAN_TU = ("=", "!=", "in", "thang", "ngay", "khoang_ngay", "chua", "khong_chua")


def _ngay_iso(s, ten="ngày"):
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(s or "")):
        raise ValueError(f"{ten} phải dạng YYYY-MM-DD")
    return datetime.date.fromisoformat(s)


def _bang(o, gt):
    """So bằng: giá trị lọc là số, hoặc ô là số và giá trị lọc đọc được số -> so số (12 = '12' = 12.0);
    còn lại so chuỗi theo `chuan`."""
    if isinstance(gt, (int, float, Decimal)) and not isinstance(gt, bool):
        a = doc_so(o)
        return a is not None and a == doc_so(gt)
    if isinstance(o, (int, float, Decimal)) and not isinstance(o, bool) and doc_so(gt) is not None:
        return doc_so(o) == doc_so(gt)
    return chuan(o) == chuan(gt)


def dung_bo_loc(loc, hmap):
    """Kiểm + biên dịch bộ lọc -> list hàm (dòng -> True/False/'NGAY_LOI')."""
    if loc in (None, ""):
        return []
    if not isinstance(loc, list):
        raise ValueError("loc phải là danh sách {cot, toan_tu, gia_tri}")
    out = []
    for i, d in enumerate(loc):
        if not isinstance(d, dict):
            raise ValueError(f"loc[{i}] phải là object")
        tt = str(d.get("toan_tu") or "=").strip()
        if tt not in TOAN_TU:
            raise ValueError(f"toan_tu '{tt}' không hỗ trợ. Có: {TOAN_TU}")
        j = tim_cot(hmap, d.get("cot"))
        gt = d.get("gia_tri")
        if tt in ("=", "!="):
            f = (lambda j, gt, neg: lambda r: (_bang(_o(r, j), gt)) != neg)(j, gt, tt == "!=")
        elif tt == "in":
            if not isinstance(gt, list) or not gt:
                raise ValueError("toan_tu 'in' cần gia_tri là danh sách")
            f = (lambda j, gt: lambda r: any(_bang(_o(r, j), g) for g in gt))(j, gt)
        elif tt in ("chua", "khong_chua"):
            k = chuan(gt)
            f = (lambda j, k, neg: lambda r: (k in chuan(_o(r, j))) != neg)(j, k, tt == "khong_chua")
        else:
            if tt == "thang":
                if not re.fullmatch(r"\d{4}-\d{2}", str(gt or "")):
                    raise ValueError("toan_tu 'thang' cần gia_tri YYYY-MM")
                y, m = map(int, gt.split("-"))
                a = datetime.date(y, m, 1)
                b = (datetime.date(y + m // 12, m % 12 + 1, 1) - datetime.timedelta(days=1))
            elif tt == "ngay":
                a = b = _ngay_iso(gt)
            else:
                if not isinstance(gt, list) or len(gt) != 2:
                    raise ValueError("toan_tu 'khoang_ngay' cần gia_tri [tu, den] (YYYY-MM-DD, gồm cả hai đầu)")
                a, b = _ngay_iso(gt[0], "tu"), _ngay_iso(gt[1], "den")

            def f(r, j=j, a=a, b=b):
                v = _o(r, j)
                dd = doc_ngay(v)
                if dd is None:
                    return "NGAY_LOI" if v not in (None, "") else False
                return a <= dd <= b
        out.append({"cot": d.get("cot"), "j": j, "toan_tu": tt, "f": f})
    return out


def _o(r, j):
    return r[j] if j < len(r) else None


# ───────────────────────────── tool: đọc ─────────────────────────────
def doc_excel(file, sheet=None, vung=None, header_dong=None, gioi_han=50, tim=None, tu_dong=None, bo_cache=False):
    """Đọc thô một sheet. Không `sheet` -> danh sách sheet. `vung` 'A1:F40' -> lưới ô có số dòng; `header_dong` -> mỗi
    dòng là {tiêu đề: giá trị}. `tim` -> chỉ dòng có ô bằng (theo `chuan`) hoặc chứa chuỗi đó. Trả tối đa `gioi_han`."""
    gioi_han = max(1, min(int(gioi_han or 50), GIOI_HAN_DONG_TRA))
    if sheet in (None, ""):
        p = duong_an_toan(file)
        return {"file": p, "sheet": _ten_sheet(p)}
    p, ten, rows = doc_sheet(file, sheet, bo_cache)
    k_tim = chuan(tim) if tim not in (None, "") else None

    def co_tim(r):
        return k_tim is None or any(chuan(c) == k_tim or (k_tim and k_tim in chuan(c)) for c in r if c is not None)

    out = {"file": p, "sheet": ten, "so_dong_sheet": len(rows)}
    if vung:
        m = re.fullmatch(r"\s*([A-Za-z]{1,3})(\d+)\s*:\s*([A-Za-z]{1,3})(\d+)\s*", str(vung))
        if not m:
            raise ValueError("vung phải dạng A1:F40")
        c1, r1, c2, r2 = so_cot(m.group(1)), int(m.group(2)), so_cot(m.group(3)), int(m.group(4))
        if r2 < r1 or c2 < c1:
            raise ValueError("vung đảo ngược")
        luoi = []
        for i in range(r1, min(r2, len(rows)) + 1):
            r = rows[i - 1]
            if not co_tim(r):
                continue
            luoi.append({"dong": i, "o": {f"{chu_cot(j)}{i}": json_gia_tri(_o(r, j)) for j in range(c1, c2 + 1)
                                         if _o(r, j) is not None}})
            if len(luoi) >= gioi_han:
                break
        out.update({"vung": vung, "dong": luoi, "cat_bot": len(luoi) >= gioi_han})
        return out
    if header_dong:
        hmap, tieu_de = header_map(rows, int(header_dong))
        inv = {j: k for k, j in hmap.items()}
        nv = {d["cot"]: d["tieu_de"] for d in tieu_de}
        ds = []
        bat_dau = max(int(tu_dong or 0), int(header_dong) + 1)
        for i in range(bat_dau, len(rows) + 1):
            r = rows[i - 1]
            if not any(c not in (None, "") for c in r) or not co_tim(r):
                continue
            ds.append({"dong": i, **{nv.get(chu_cot(j), chu_cot(j)) if j in inv else chu_cot(j): json_gia_tri(c)
                                     for j, c in enumerate(r) if c is not None}})
            if len(ds) >= gioi_han:
                break
        out.update({"header_dong": int(header_dong), "tieu_de": tieu_de, "dong": ds, "cat_bot": len(ds) >= gioi_han})
        return out
    ds = []
    for i in range(max(1, int(tu_dong or 1)), len(rows) + 1):
        r = rows[i - 1]
        if not any(c not in (None, "") for c in r) or not co_tim(r):
            continue
        ds.append({"dong": i, "o": {chu_cot(j): json_gia_tri(c) for j, c in enumerate(r) if c is not None}})
        if len(ds) >= gioi_han:
            break
    out.update({"dong": ds, "cat_bot": len(ds) >= gioi_han})
    return out


# ───────────────────────────── tool: tổng hợp ─────────────────────────────
def tong_hop_excel(file, sheet, header_dong, loc=None, cot_cong=None, nhom_theo=None, tu_dong=None, den_dong=None,
                   bo_cache=False):
    """Σ chính xác các cột `cot_cong` trên các dòng (dưới header) khớp MỌI điều kiện `loc` (AND). Có `nhom_theo` (tên
    cột hoặc danh sách) thì tách tổng theo nhóm. Trả số dòng khớp, ≤40 dòng góp phần (số dòng Excel + ô lọc + ô cộng),
    và thống kê ô bị bỏ (ngày không đọc được, ô cộng không phải số)."""
    if not cot_cong:
        raise ValueError("cần cot_cong (danh sách cột để cộng)")
    if isinstance(cot_cong, str):
        cot_cong = [cot_cong]
    p, ten, rows = doc_sheet(file, sheet, bo_cache)
    hmap, tieu_de = header_map(rows, int(header_dong))
    bl = dung_bo_loc(loc, hmap)
    jc = [(c, tim_cot(hmap, c)) for c in cot_cong]
    nhom = [] if nhom_theo in (None, "", []) else ([nhom_theo] if isinstance(nhom_theo, str) else list(nhom_theo))
    jn = [(c, tim_cot(hmap, c)) for c in nhom]
    tong = {c: Decimal(0) for c, _ in jc}
    theo_nhom = {}
    gop, so_dong = [], 0
    ngay_loi, khong_so = {}, {}
    a = int(header_dong) + 1 if tu_dong is None else max(int(tu_dong), int(header_dong) + 1)
    b = len(rows) if den_dong is None else min(int(den_dong), len(rows))
    for i in range(a, b + 1):
        r = rows[i - 1]
        if not any(c not in (None, "") for c in r):
            continue
        ok = True
        for d in bl:
            kq = d["f"](r)
            if kq == "NGAY_LOI":
                lst = ngay_loi.setdefault(d["cot"], [])
                if len(lst) < 5:
                    lst.append({"dong": i, "o": json_gia_tri(_o(r, d["j"]))})
                else:
                    lst.append(None)
                ok = False
                break
            if not kq:
                ok = False
                break
        if not ok:
            continue
        so_dong += 1
        k_nhom = tuple(json_gia_tri(_o(r, j)) for _, j in jn) if jn else None
        if jn:
            g = theo_nhom.setdefault(k_nhom, {"so_dong": 0, **{c: Decimal(0) for c, _ in jc}})
            g["so_dong"] += 1
        for c, j in jc:
            v = doc_so(_o(r, j))
            if v is None:
                if _o(r, j) not in (None, ""):
                    lst = khong_so.setdefault(c, [])
                    lst.append({"dong": i, "o": json_gia_tri(_o(r, j))} if len(lst) < 5 else None)
                continue
            tong[c] += v
            if jn:
                g[c] += v
        if len(gop) < GIOI_HAN_DONG_GOP:
            o = {f"{x['cot']}": json_gia_tri(_o(r, x["j"])) for x in bl}
            o.update({c: json_gia_tri(_o(r, j)) for c, j in jc})
            o.update({c: json_gia_tri(_o(r, j)) for c, j in jn})
            gop.append({"dong": i, **o})

    def gon(d):
        return {k: {"so_o_loi": len(v), "mau": [x for x in v if x]} for k, v in d.items()}
    out = {"file": p, "sheet": ten, "header_dong": int(header_dong),
           "loc": [{"cot": d["cot"], "cot_excel": chu_cot(d["j"]), "toan_tu": d["toan_tu"]} for d in bl],
           "so_dong_khop": so_dong,
           "tong": {c: _so_ra(v) for c, v in tong.items()},
           "tong_chinh_xac": {c: str(v) for c, v in tong.items()},
           "dong_gop_phan": gop, "cat_bot_dong_gop": max(0, so_dong - len(gop)),
           "o_ngay_khong_doc_duoc": gon(ngay_loi), "o_khong_phai_so": gon(khong_so),
           "quy_uoc": "so chuỗi: NFC + bỏ dấu cách thừa + không phân biệt hoa thường (giữ dấu); ngày: datetime/serial "
                      "Excel/dd/mm/yyyy (ngày trước tháng); cộng Decimal không làm tròn"}
    if jn:
        ds = sorted(theo_nhom.items(), key=lambda kv: -abs(kv[1][jc[0][0]]))
        out["theo_nhom"] = [{"nhom": dict(zip([c for c, _ in jn], k)), "so_dong": g["so_dong"],
                             **{c: _so_ra(g[c]) for c, _ in jc}} for k, g in ds[:200]]
        out["so_nhom"] = len(ds)
    return out
