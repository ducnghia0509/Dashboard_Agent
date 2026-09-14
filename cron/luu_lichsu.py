# -*- coding: utf-8 -*-
"""LƯU LỊCH SỬ NGUỒN theo NGÀY cho khối Dòng tiền.

VÌ SAO CẦN: nguồn số dư / báo cáo ngân hàng bị GHI ĐÈ mỗi lượt kéo (receiver ghi "wb"), trong khi
công thức mapping là DIFF "báo cáo ngày N trừ báo cáo ngày N-1". Không giữ lịch sử thì bản ngày
N-1 mất sạch, chỉ còn cách cất số vào payload DB — không tra ngược được, hỏng 1 lượt cron là đứt
chuỗi. Ở đây mỗi lượt kéo được lưu lại nguyên vẹn để diff đọc thẳng từ file thật.

ĐỂ NGOÀI `received_reports/` LÀ CỐ Ý: `cashflow_vay_extractor.files_can_doc()` quét ĐỆ QUY cả
`THUCHI/` rồi chọn bản mtime MỚI NHẤT cho mỗi (đơn vị, tháng). Đặt bản lưu trữ trong đó là nó
tranh chỗ với file nguồn chính và âm thầm đổi số của màn THÁNG.

Hai kiểu lưu, theo hình dạng nguồn:
  · SỐ DƯ (mỗi file 1 sheet chính) -> "thêm 1 sheet theo ngày": dồn vào 1 workbook lưu trữ cho mỗi
    (đơn vị, ngân hàng), mỗi lượt kéo là 1 sheet tên `YYYY-MM-DD HHh`. Số dư khả dụng kéo 3
    lượt/ngày (9h/13h/18h) nên tên sheet PHẢI có giờ, không thì lượt sau đè lượt trước.
  · BÁO CÁO NGÂN HÀNG (7+ sheet: DATA, BCTH 2, NH VPB, TH VPB, TH BAB…) -> không "thêm 1 sheet"
    được, lưu NGUYÊN FILE vào thư mục theo ngày `<kho>/baocaonganhang/<YYYY-MM-DD>/<tên gốc>`.
    Mapping cũng viết "báo cáo ngày N", không viết "sheet ngày N".
"""
import os
import shutil
from datetime import datetime, timedelta, timezone

import openpyxl

VN = timezone(timedelta(hours=7))
CONNECT = "/home/itadmin/AI_Dashboard_QT/Connect_VPS"
KHO = os.path.join(CONNECT, "lichsu_nguon")
KHO_SODU = os.path.join(KHO, "sodu")
KHO_BCNH = os.path.join(KHO, "baocaonganhang")

_CAM = set('[]:*?/\\')          # ký tự Excel cấm trong tên sheet


def ten_sheet(dt=None) -> str:
    """`2026-09-14 09h` — có GIỜ vì số dư kéo 3 lượt/ngày. Tối đa 31 ký tự (giới hạn Excel)."""
    dt = dt or datetime.now(VN)
    return "".join(c for c in f"{dt:%Y-%m-%d %Hh}" if c not in _CAM)[:31]


def ten_ngay(dt=None) -> str:
    return f"{(dt or datetime.now(VN)):%Y-%m-%d}"


def _doc_gia_tri(path: str) -> list:
    """Ma trận giá trị của sheet ĐẦU TIÊN (sheet extractor đang đọc). .xls qua xlrd."""
    if path.lower().endswith(".xls"):
        import xlrd
        ws = xlrd.open_workbook(path).sheet_by_index(0)
        return [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)]
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def luu_sheet_sodu(src: str, nhan: str = None, log=print) -> str:
    """Chép sheet chính của `src` vào workbook lưu trữ cùng tên, đặt tên sheet theo ngày+giờ.

    Trùng tên sheet (chạy lại trong cùng khung giờ) -> GHI ĐÈ sheet đó, đúng tinh thần "giữ bản
    mới nhất của mỗi mốc" mà cron đang dùng cho raw_rows.
    """
    nhan = nhan or ten_sheet()
    os.makedirs(KHO_SODU, exist_ok=True)
    dest = os.path.join(KHO_SODU, os.path.splitext(os.path.basename(src))[0] + ".xlsx")
    try:
        rows = _doc_gia_tri(src)
    except Exception as ex:                                    # noqa: BLE001
        log(f"  LƯU TRỮ LỖI đọc {os.path.basename(src)}: {ex}")
        return ""
    wb = openpyxl.load_workbook(dest) if os.path.exists(dest) else openpyxl.Workbook()
    if not os.path.exists(dest):
        wb.remove(wb.active)                                   # bỏ sheet rỗng mặc định
    if nhan in wb.sheetnames:
        wb.remove(wb[nhan])
    ws = wb.create_sheet(title=nhan)
    for r in rows:
        ws.append(list(r))
    # Excel giới hạn 255 sheet/workbook trên vài phiên bản cũ; 3 lượt/ngày ~ 90 sheet/tháng nên
    # cắt bớt sheet CŨ NHẤT khi vượt 200 (giữ ~2 tháng gần nhất, đủ cho mọi diff ngày).
    while len(wb.sheetnames) > 200:
        wb.remove(wb[sorted(wb.sheetnames)[0]])
    wb.save(dest)
    wb.close()
    return dest


def luu_file_bcnh(src: str, ngay: str = None, log=print) -> str:
    """Lưu NGUYÊN file báo cáo ngân hàng vào `<kho>/baocaonganhang/<ngày>/<tên gốc>`."""
    ngay = ngay or ten_ngay()
    thu_muc = os.path.join(KHO_BCNH, ngay)
    os.makedirs(thu_muc, exist_ok=True)
    dest = os.path.join(thu_muc, os.path.basename(src))
    try:
        shutil.copy2(src, dest)
    except OSError as ex:
        log(f"  LƯU TRỮ LỖI chép {os.path.basename(src)}: {ex}")
        return ""
    return dest


def file_bcnh_cua_ngay(ngay: str) -> list:
    """Danh sách file báo cáo ngân hàng đã lưu của đúng ngày đó ([] nếu chưa có)."""
    thu_muc = os.path.join(KHO_BCNH, ngay)
    if not os.path.isdir(thu_muc):
        return []
    return sorted(os.path.join(thu_muc, f) for f in os.listdir(thu_muc)
                  if f.lower().endswith((".xlsx", ".xlsm", ".xls")) and not f.startswith("~$"))


def ngay_bcnh_co_san() -> list:
    """Các ngày đã có bản lưu báo cáo ngân hàng, cũ -> mới."""
    if not os.path.isdir(KHO_BCNH):
        return []
    return sorted(d for d in os.listdir(KHO_BCNH) if os.path.isdir(os.path.join(KHO_BCNH, d)))
